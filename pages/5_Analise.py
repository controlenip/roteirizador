
import io
import re
import time
from concurrent.futures import ThreadPoolExecutor
import html
import hashlib
import zipfile
import unicodedata
import xml.etree.ElementTree as ET

from collections import defaultdict
from datetime import datetime, timedelta

import numpy as np
import pandas as pd
import requests
import streamlit as st
import folium

from folium.plugins import MarkerCluster
from streamlit_folium import st_folium
from sklearn.cluster import DBSCAN
from sklearn.neighbors import BallTree
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# CONFIGURACAO GERAL
# ============================================================

st.set_page_config(
    page_title="NIP | Análise Cruzada",
    page_icon="📍",
    layout="wide",
    initial_sidebar_state="expanded"
)

CAT_SAN = "OBRA SANEAMENTO"
CAT_LEV = "OBRA LEVANTAMENTO"
CAT_DUP = "OBRA SANEAMENTO E LEVANTAMENTO"

CATS = [
    CAT_SAN,
    CAT_LEV,
    CAT_DUP
]

COLORS = {
    CAT_SAN: "#2563eb",
    CAT_LEV: "#16a34a",
    CAT_DUP: "#9333ea"
}

MAP_COLORS = {
    CAT_SAN: "blue",
    CAT_LEV: "green",
    CAT_DUP: "purple"
}

LIST_VALID = {
    "0",
    "EM LEVANTAMENTO",
    "CORRECAO DE LEVANTAMENTO"
}

CONTRACT_VALID = {
    "0",
    "NIP GLOBAL LTDA - EQTL MARANHAO"
}

R = 6371.0088

OSRM = "https://router.project-osrm.org"

DAY_NAMES = [
    "Segunda",
    "Terça",
    "Quarta",
    "Quinta",
    "Sexta",
    "Sábado",
    "Domingo"
]


# ============================================================
# ESTILO VISUAL
# ============================================================

st.markdown(
    """
    <style>
    .block-container {
        padding-top: 4.5rem !important;
        padding-bottom: 2rem !important;
        max-width: 100% !important;
    }

    .nip-title {
        font-size: 2rem;
        color: #0d256c;
        font-weight: 800;
        line-height: 1.5;
        margin: 0 0 8px;
        overflow: visible;
        padding-top: 8px;
    }

    .nip-muted {
        color: #64748b;
        margin-bottom: 20px;
    }

    .nip-card {
        padding: 15px;
        border: 1px solid #e2e8f0;
        border-left: 5px solid var(--accent);
        border-radius: 10px;
        min-height: 110px;
        overflow-wrap: anywhere;
    }

    .nip-label {
        font-size: 12px;
        color: #64748b;
        font-weight: 600;
    }

    .nip-value {
        font-size: 26px;
        font-weight: 800;
        color: #172554;
    }

    .nip-small {
        font-size: 11px;
        color: #64748b;
    }

    @media(max-width:900px) {
        .block-container {
            padding-top: 3.3rem !important;
        }

        .nip-title {
            font-size: 1.45rem;
        }
    }
    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# NORMALIZACAO
# ============================================================

def n(v):
    if v is None:
        return ""

    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass

    texto = unicodedata.normalize(
        "NFKD",
        str(v)
    ).encode(
        "ascii",
        "ignore"
    ).decode().upper().strip()

    return re.sub(r"\s+", " ", texto)


def clean(v):
    s = n(v)

    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]

    return s


def nota(v):
    s = str(v).strip() if pd.notna(v) else ""

    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]

    if n(s) in (
        "",
        "NAN",
        "NONE",
        "NULL",
        "0"
    ):
        return ""

    return s


def col(df, variants, required=True):
    names = {
        n(c): c
        for c in df.columns
    }

    for x in variants:
        if n(x) in names:
            return names[n(x)]

    if required:
        raise ValueError(
            "Coluna ausente: "
            + " / ".join(variants)
        )

    return None


def to_num(series):
    return pd.to_numeric(
        series.astype(str).str.replace(
            ",",
            ".",
            regex=False
        ),
        errors="coerce"
    )


def valid_coords(df, lat, lon):
    la = to_num(df[lat])
    lo = to_num(df[lon])

    ok = (
        la.between(-35, 6)
        & lo.between(-75, -30)
        & la.ne(0)
        & lo.ne(0)
    )

    return la.where(ok), lo.where(ok)


# ============================================================
# IDENTIFICACAO AUTOMATICA DE ABAS
# ============================================================

def find_sheet(file, groups, prefer=""):
    xls = pd.ExcelFile(
        io.BytesIO(file.getvalue()),
        engine="openpyxl"
    )

    names = sorted(
        xls.sheet_names,
        key=lambda s: (
            0 if n(s) == n(prefer) else 1
        )
    )

    for name in names:
        for header in range(6):
            try:
                test = pd.read_excel(
                    xls,
                    sheet_name=name,
                    header=header,
                    nrows=3,
                    dtype=str
                )

                keys = {
                    n(x)
                    for x in test.columns
                }

                encontrado = all(
                    any(
                        n(v) in keys
                        for v in group
                    )
                    for group in groups
                )

                if encontrado:
                    df = pd.read_excel(
                        xls,
                        sheet_name=name,
                        header=header,
                        dtype=str
                    )

                    return df, name

            except (
                ValueError,
                TypeError,
                KeyError
            ):
                continue

    raise ValueError(
        "Não encontrei aba com as colunas exigidas. "
        "Abas: "
        + ", ".join(xls.sheet_names)
    )


# ============================================================
# DATAS
# ============================================================

def parse_date(v):
    if v is None or n(v) in (
        "",
        "NAN",
        "NONE",
        "NULL",
        "0",
        "-"
    ):
        return pd.NaT

    try:
        f = float(
            str(v).replace(",", ".")
        )

        if 20000 <= f <= 80000:
            return (
                pd.Timestamp("1899-12-30")
                + pd.Timedelta(days=f)
            )

    except (
        ValueError,
        TypeError,
        OverflowError
    ):
        pass

    return pd.to_datetime(
        v,
        dayfirst=True,
        errors="coerce"
    )


# ============================================================
# LEITURA SANEAMENTO
# ============================================================

def load_san(file):
    if file.name.lower().endswith(".csv"):
        raw = file.getvalue()

        try:
            df = pd.read_csv(
                io.BytesIO(raw),
                sep=None,
                engine="python",
                dtype=str,
                encoding="utf-8-sig"
            )

        except UnicodeDecodeError:
            df = pd.read_csv(
                io.BytesIO(raw),
                sep=None,
                engine="python",
                dtype=str,
                encoding="latin-1"
            )

        sheet = "CSV"

    else:
        df, sheet = find_sheet(
            file,
            [
                ["NOTA"],
                [
                    "MUNICIPIO",
                    "MUNICÍPIO",
                    "CIDADE"
                ]
            ],
            "Clientes Existentes"
        )

    k = col(df, ["NOTA"])

    city = col(
        df,
        [
            "MUNICIPIO",
            "MUNICÍPIO",
            "CIDADE"
        ]
    )

    reg = col(
        df,
        ["REGIONAL"],
        False
    )

    lat = col(
        df,
        [
            "LATITUDE PROJETO",
            "LATITUDE",
            "LATITUDE CAMPO"
        ]
    )

    lon = col(
        df,
        [
            "LONGITUDE PROJETO",
            "LONGITUDE",
            "LONGITUDE CAMPO"
        ]
    )

    df["CHAVE"] = df[k].map(nota)

    df["CIDADE_OBRA"] = (
        df[city].fillna("")
    )

    df["REGIONAL_OBRA"] = (
        df[reg].fillna("")
        if reg else ""
    )

    df["LAT_OBRA"], df["LON_OBRA"] = (
        valid_coords(
            df,
            lat,
            lon
        )
    )

    return df, sheet


# ============================================================
# LEITURA LEVANTAMENTO
# ============================================================

def load_lev(file):
    df, sheet = find_sheet(
        file,
        [
            ["PROTOCOLO"],
            [
                "STATUS SAP",
                "STATUS_SAP"
            ],
            [
                "STATUS LIST",
                "STATUS_LIST"
            ]
        ],
        "NOTAS"
    )

    k = col(df, ["PROTOCOLO"])

    sap = col(
        df,
        [
            "STATUS SAP",
            "STATUS_SAP"
        ]
    )

    lst = col(
        df,
        [
            "STATUS LIST",
            "STATUS_LIST"
        ]
    )

    ctr = col(
        df,
        ["CONTRATO"]
    )

    city = col(
        df,
        [
            "MUNICIPIO",
            "MUNICÍPIO",
            "CIDADE"
        ]
    )

    reg = col(
        df,
        ["REGIONAL"],
        False
    )

    lat = col(df, ["LATITUDE"])
    lon = col(df, ["LONGITUDE"])

    m = next(
        (
            c for c in df.columns
            if "ORCAMENTO MODULAR" in n(c)
        ),
        None
    )

    nn = next(
        (
            c for c in df.columns
            if "PLA ALVOS" in n(c)
        ),
        None
    )

    if m is None or nn is None:
        raise ValueError(
            "Não localizei ORÇAMENTO MODULAR "
            "e PLA ALVOS."
        )

    if len(df.columns) < 3:
        raise ValueError(
            "A aba necessita das colunas físicas B e C."
        )

    b = df.columns[1]
    c = df.columns[2]

    priority = col(
        df,
        ["PRIORIDADE"],
        False
    )

    opening = col(
        df,
        [
            "DATA ABERTURA",
            "DATA DE ABERTURA"
        ],
        False
    )

    df["CHAVE"] = df[k].map(nota)

    df["CIDADE_OBRA"] = (
        df[city].fillna("")
    )

    df["REGIONAL_OBRA"] = (
        df[reg].fillna("")
        if reg else ""
    )

    df["LAT_OBRA"], df["LON_OBRA"] = (
        valid_coords(
            df,
            lat,
            lon
        )
    )

    df["SAP"] = df[sap].map(clean)
    df["LIST"] = df[lst].map(clean)

    df["CONTRATO_N"] = (
        df[ctr].map(clean)
    )

    df["M"] = df[m].map(clean)
    df["N"] = df[nn].map(clean)

    df["RESP_CAMPO"] = (
        df[b]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    df["DATA_CAMPO_DT"] = (
        df[c].map(parse_date)
    )

    df["EM_CAMPO"] = (
        df["RESP_CAMPO"].map(
            lambda x: n(x) not in (
                "",
                "SEM LEVANTADOR",
                "NAN",
                "NONE",
                "NULL",
                "0",
                "-"
            )
        )
        & df["DATA_CAMPO_DT"].notna()
    )

    df["PRIORIDADE_RAW"] = (
        df[priority].fillna("")
        if priority else ""
    )

    df["ABERTURA_RAW"] = (
        df[opening].fillna("")
        if opening else ""
    )

    def reasons(r):
        result = []

        if r["SAP"] in (
            "FINL",
            "CANC"
        ):
            result.append(
                "SAP " + r["SAP"]
            )

        if r["LIST"] not in LIST_VALID:
            result.append(
                "LIST "
                + (
                    r["LIST"]
                    or "VAZIO"
                )
            )

        if r["CONTRATO_N"] not in CONTRACT_VALID:
            result.append(
                "CONTRATO NÃO PERMITIDO"
            )

        if r["M"] != "0":
            result.append(
                "ORÇAMENTO MODULAR NÃO ZERO"
            )

        if r["N"] != "0":
            result.append(
                "PLA ALVOS NÃO ZERO"
            )

        if r["EM_CAMPO"]:
            result.append(
                "JÁ EM CAMPO (B+C)"
            )

        return result

    df["MOTIVOS"] = df.apply(
        reasons,
        axis=1
    )

    return df, {
        "aba": sheet,
        "B": b,
        "C": c,
        "linhas_campo": int(
            df["EM_CAMPO"].sum()
        )
    }


# ============================================================
# LEITURA EQUIPES
# ============================================================

def load_teams(file):
    xls = pd.ExcelFile(
        io.BytesIO(file.getvalue()),
        engine="openpyxl"
    )

    parts = []

    for kind, sheet in [
        (
            "SANEAMENTO",
            "SANEAMENTO"
        ),
        (
            "LEVANTAMENTO",
            "LEVANTADORES"
        )
    ]:
        name = next(
            (
                x for x in xls.sheet_names
                if n(x) == sheet
            ),
            None
        )

        if name is None:
            raise ValueError(
                "Aba de equipe ausente: "
                + sheet
            )

        df = pd.read_excel(
            xls,
            sheet_name=name,
            dtype=str
        )

        namecol = col(
            df,
            [
                "NOME",
                "EQUIPE",
                "NOME_COLAB"
            ]
        )

        city = col(
            df,
            [
                "CIDADES",
                "CIDADE",
                "MUNICIPIO"
            ]
        )

        lat = col(
            df,
            [
                "LATITUDE",
                "LAT"
            ]
        )

        lon = col(
            df,
            [
                "LONGITUDE",
                "LON"
            ]
        )

        la, lo = valid_coords(
            df,
            lat,
            lon
        )

        parts.append(
            pd.DataFrame({
                "TIPO_EQUIPE": kind,
                "EQUIPE": df[namecol],
                "CIDADE_BASE": df[city],
                "LAT_EQUIPE": la,
                "LON_EQUIPE": lo
            })
        )

    data = pd.concat(
        parts,
        ignore_index=True
    ).dropna(
        subset=[
            "EQUIPE",
            "LAT_EQUIPE",
            "LON_EQUIPE"
        ]
    )

    data = data[
        data["EQUIPE"]
        .astype(str)
        .str.strip()
        .ne("")
    ].copy()

    data["EQUIPE"] = (
        data["EQUIPE"]
        .astype(str)
        .str.strip()
    )

    data["CIDADE_BASE"] = (
        data["CIDADE_BASE"].fillna("")
    )

    data["ID_EQUIPE"] = (
        data["TIPO_EQUIPE"]
        + "|"
        + data["EQUIPE"]
        + "|"
        + data["CIDADE_BASE"]
    )

    return data.drop_duplicates(
        "ID_EQUIPE"
    ).reset_index(drop=True)


# ============================================================
# CRUZAMENTO DAS BASES
# ============================================================

def consolidate(san, lev):
    gs = dict(
        tuple(
            san[
                san["CHAVE"].ne("")
            ].groupby(
                "CHAVE",
                sort=False
            )
        )
    )

    gl = dict(
        tuple(
            lev[
                lev["CHAVE"].ne("")
            ].groupby(
                "CHAVE",
                sort=False
            )
        )
    )

    rows = []

    for key in dict.fromkeys(
        [*gs, *gl]
    ):
        a = gs.get(key)
        b = gl.get(key)

        if a is not None:
            valid_a = (
                a["LAT_OBRA"].notna()
                & a["LON_OBRA"].notna()
            )

            sa = (
                a.loc[valid_a].iloc[0]
                if valid_a.any()
                else a.iloc[0]
            )
        else:
            sa = None

        if b is not None:
            valid_b = (
                b["LAT_OBRA"].notna()
                & b["LON_OBRA"].notna()
            )

            sb = (
                b.loc[valid_b].iloc[0]
                if valid_b.any()
                else b.iloc[0]
            )
        else:
            sb = None

        source = (
            sb
            if (
                sb is not None
                and pd.notna(
                    sb["LAT_OBRA"]
                )
            )
            else sa
            if sa is not None
            else sb
        )

        reasons = (
            list(
                dict.fromkeys(
                    v
                    for sub in b["MOTIVOS"]
                    for v in sub
                )
            )
            if b is not None
            else []
        )

        status_saps = (
            set(b["SAP"])
            if b is not None
            else set()
        )

        on_field = (
            b[b["EM_CAMPO"]]
            if b is not None
            else pd.DataFrame()
        )

        rows.append({
            "NOTA": key,
            "LISTA": (
                CAT_DUP
                if (
                    a is not None
                    and b is not None
                )
                else CAT_SAN
                if a is not None
                else CAT_LEV
            ),
            "MUNICIPIO": source[
                "CIDADE_OBRA"
            ],
            "REGIONAL": source[
                "REGIONAL_OBRA"
            ],
            "LATITUDE": source[
                "LAT_OBRA"
            ],
            "LONGITUDE": source[
                "LON_OBRA"
            ],
            "DUPLICADA": (
                "SIM"
                if (
                    a is not None
                    and b is not None
                )
                else "NÃO"
            ),
            "PENDENTE_CONTAGEM": (
                "NÃO"
                if reasons
                else "SIM"
            ),
            "MOTIVO_EXCLUSAO": (
                " | ".join(reasons)
                if reasons
                else "-"
            ),
            "SAP_FINL": (
                "SIM"
                if "FINL" in status_saps
                else "NÃO"
            ),
            "SAP_CANC": (
                "SIM"
                if "CANC" in status_saps
                else "NÃO"
            ),
            "JA_EM_CAMPO": (
                "SIM"
                if len(on_field)
                else "NÃO"
            ),
            "EQUIPE_CAMPO": (
                " | ".join(
                    on_field[
                        "RESP_CAMPO"
                    ].astype(str).unique()
                )
                if len(on_field)
                else ""
            ),
            "DATA_CAMPO": (
                " | ".join(
                    pd.Timestamp(x).strftime(
                        "%d/%m/%Y"
                    )
                    for x in on_field[
                        "DATA_CAMPO_DT"
                    ].dropna().unique()
                )
                if len(on_field)
                else ""
            ),
            "STATUS_SAP": (
                " | ".join(
                    dict.fromkeys(
                        b["SAP"]
                    )
                )
                if b is not None
                else ""
            ),
            "STATUS_LIST": (
                " | ".join(
                    dict.fromkeys(
                        b["LIST"]
                    )
                )
                if b is not None
                else ""
            ),
            "PRIORIDADE_ORIGINAL": (
                sb["PRIORIDADE_RAW"]
                if sb is not None
                else ""
            ),
            "DATA_ABERTURA": (
                sb["ABERTURA_RAW"]
                if sb is not None
                else ""
            ),
            "OCORRENCIAS_SANEAMENTO": (
                len(a)
                if a is not None
                else 0
            ),
            "OCORRENCIAS_LEVANTAMENTO": (
                len(b)
                if b is not None
                else 0
            ),
            "MUNICIPIO_SANEAMENTO": (
                sa["CIDADE_OBRA"]
                if sa is not None
                else ""
            ),
            "MUNICIPIO_LEVANTAMENTO": (
                sb["CIDADE_OBRA"]
                if sb is not None
                else ""
            )
        })

    return pd.DataFrame(rows)


# ============================================================
# TIPOS DE ATIVIDADE
# ============================================================

def activity(cat):
    if cat == CAT_DUP:
        return [
            "SANEAMENTO",
            "LEVANTAMENTO"
        ]

    if cat == CAT_SAN:
        return ["SANEAMENTO"]

    return ["LEVANTAMENTO"]


# ============================================================
# HAVERSINE
# ============================================================

def haversine(a, b, c, d):
    x = np.radians(a)
    y = np.radians(b)
    z = np.radians(c)
    w = np.radians(d)

    t = (
        np.sin(
            (z - x) / 2
        ) ** 2
        + np.cos(x)
        * np.cos(z)
        * np.sin(
            (w - y) / 2
        ) ** 2
    )

    return (
        2
        * R
        * np.arcsin(
            np.sqrt(
                np.clip(t, 0, 1)
            )
        )
    )


# ============================================================
# PRIORIDADE
# ============================================================

def priority(df, days):
    df = df.copy()

    dates = pd.to_datetime(
        df["DATA_ABERTURA"],
        errors="coerce",
        dayfirst=True
    )

    age = (
        pd.Timestamp.today().normalize()
        - dates
    ).dt.days

    df["DIAS_ABERTURA"] = age

    resultados = []

    for p, d in zip(
        df["PRIORIDADE_ORIGINAL"],
        age
    ):
        if (
            "ALTA" in n(p)
            or "URGENTE" in n(p)
            or (
                pd.notna(d)
                and d >= 4 * days
            )
        ):
            resultados.append(
                "ALTA"
            )

        elif (
            pd.notna(d)
            and d >= days
        ):
            resultados.append(
                "MEDIA"
            )

        else:
            resultados.append(
                "BAIXA"
            )

    df["PRIORIDADE_PLANEJAMENTO"] = (
        resultados
    )

    return df


# ============================================================
# CANDIDATOS MAIS PROXIMOS
# ============================================================

@st.cache_data(
    show_spinner=False
)
def team_candidates(
    obras,
    equipes,
    k
):
    rows = []
    labels = {}

    for kind in [
        "SANEAMENTO",
        "LEVANTAMENTO"
    ]:
        team = equipes[
            equipes[
                "TIPO_EQUIPE"
            ].eq(kind)
        ].reset_index(drop=True)

        eligible = obras[
            obras["LISTA"].map(
                lambda c: (
                    kind in activity(c)
                )
            )
        ].dropna(
            subset=[
                "LATITUDE",
                "LONGITUDE"
            ]
        )

        if team.empty or eligible.empty:
            continue

        tree = BallTree(
            np.radians(
                team[
                    [
                        "LAT_EQUIPE",
                        "LON_EQUIPE"
                    ]
                ].to_numpy(float)
            ),
            metric="haversine"
        )

        dist, ids = tree.query(
            np.radians(
                eligible[
                    [
                        "LATITUDE",
                        "LONGITUDE"
                    ]
                ].to_numpy(float)
            ),
            k=min(
                k,
                len(team)
            )
        )

        for i, (_, work) in enumerate(
            eligible.iterrows()
        ):
            txt = []

            for j in range(
                ids.shape[1]
            ):
                eq = team.iloc[
                    ids[i, j]
                ]

                km = round(
                    float(
                        dist[i, j] * R
                    ),
                    2
                )

                rows.append({
                    "NOTA": work["NOTA"],
                    "ATIVIDADE": kind,
                    "ID_EQUIPE": eq["ID_EQUIPE"],
                    "EQUIPE": eq["EQUIPE"],
                    "CIDADE_BASE": eq["CIDADE_BASE"],
                    "LAT_EQUIPE": eq["LAT_EQUIPE"],
                    "LON_EQUIPE": eq["LON_EQUIPE"],
                    "LAT_OBRA": work["LATITUDE"],
                    "LON_OBRA": work["LONGITUDE"],
                    "KM_RETA": km
                })

                txt.append(
                    f"{eq['EQUIPE']} "
                    f"({eq['CIDADE_BASE']}) "
                    f"- {km:.1f} km (linha reta)"
                )

            labels[
                (
                    work["NOTA"],
                    kind
                )
            ] = " | ".join(txt)

    out = obras.copy()

    for kind in [
        "SANEAMENTO",
        "LEVANTAMENTO"
    ]:
        out[
            "EQUIPES_" + kind
        ] = out.apply(
            lambda r: labels.get(
                (
                    r["NOTA"],
                    kind
                ),
                (
                    "SEM COORDENADAS OU EQUIPE"
                    if kind in activity(
                        r["LISTA"]
                    )
                    else "NÃO APLICÁVEL"
                )
            ),
            axis=1
        )

    columns = [
        "NOTA",
        "ATIVIDADE",
        "ID_EQUIPE",
        "EQUIPE",
        "CIDADE_BASE",
        "LAT_EQUIPE",
        "LON_EQUIPE",
        "LAT_OBRA",
        "LON_OBRA",
        "KM_RETA"
    ]

    return out, pd.DataFrame(
        rows,
        columns=columns
    )


# ============================================================
# OSRM COM CACHE E TIMEOUT
# ============================================================

@st.cache_data(
    ttl=86400,
    max_entries=10000,
    show_spinner=False
)
def osrm_route(
    lat1,
    lon1,
    lat2,
    lon2,
    server,
    geometry=False
):
    if (
        lat1 == lat2
        and lon1 == lon2
    ):
        return {
            "ok": True,
            "km": 0.0,
            "min": 0.0,
            "geometry": [
                [lon1, lat1],
                [lon2, lat2]
            ]
        }

    url = (
        server.rstrip("/")
        + "/route/v1/driving/"
        + f"{lon1:.6f},{lat1:.6f};"
        + f"{lon2:.6f},{lat2:.6f}"
    )

    try:
        response = requests.get(
            url,
            params={
                "overview": (
                    "full"
                    if geometry
                    else "false"
                ),
                "geometries": "geojson",
                "steps": "false"
            },
            headers={
                "User-Agent": "NIP-Planejamento/1.0"
            },
            timeout=(2, 4)
        )

        response.raise_for_status()

        data = response.json()

        if (
            data.get("code") != "Ok"
            or not data.get("routes")
        ):
            return {
                "ok": False
            }

        route = data["routes"][0]

        geom = (
            route.get(
                "geometry",
                {}
            ).get(
                "coordinates",
                []
            )
            if geometry
            else []
        )

        return {
            "ok": True,
            "km": round(
                float(
                    route["distance"]
                ) / 1000,
                2
            ),
            "min": round(
                float(
                    route["duration"]
                ) / 60,
                1
            ),
            "geometry": geom
        }

    except (
        requests.RequestException,
        ValueError,
        TypeError,
        KeyError
    ):
        return {
            "ok": False
        }


# ============================================================
# PREPARACAO DE ROTAS
# ============================================================

def initial_routes(
    candidates,
    max_km
):
    data = candidates.copy()

    data["KM_RODOVIARIO"] = np.nan
    data["TEMPO_RODOVIARIO_MIN"] = np.nan

    data["ROTA_STATUS"] = (
        "NÃO CONSULTADA"
    )

    data["KM_ALOCACAO"] = np.nan

    data.loc[
        data["KM_RETA"].gt(max_km),
        "ROTA_STATUS"
    ] = "FORA RAIO GEOGRÁFICO"

    return data


# ============================================================
# CONSULTAS OSRM POR LOTES
# ============================================================

def compute_route_batch(
    routes,
    start,
    batch,
    server,
    placeholder=None
):
    df = routes.copy()

    eligible = df.index[
        df["ROTA_STATUS"].eq(
            "NÃO CONSULTADA"
        )
    ].tolist()

    end = min(
        start + batch,
        len(eligible)
    )

    t0 = time.monotonic()

    for pos in range(start, end):
        idx = eligible[pos]

        r = df.loc[idx]

        answer = osrm_route(
            round(
                float(
                    r["LAT_EQUIPE"]
                ),
                6
            ),
            round(
                float(
                    r["LON_EQUIPE"]
                ),
                6
            ),
            round(
                float(
                    r["LAT_OBRA"]
                ),
                6
            ),
            round(
                float(
                    r["LON_OBRA"]
                ),
                6
            ),
            server,
            False
        )

        if answer["ok"]:
            df.at[
                idx,
                "KM_RODOVIARIO"
            ] = answer["km"]

            df.at[
                idx,
                "TEMPO_RODOVIARIO_MIN"
            ] = answer["min"]

            df.at[
                idx,
                "ROTA_STATUS"
            ] = "OSRM CONFIRMADA"

        else:
            df.at[
                idx,
                "ROTA_STATUS"
            ] = "SEM ROTA OSRM"

        if placeholder is not None:
            placeholder.text(
                f"Lote: "
                f"{pos - start + 1}/{end - start}"
                f" | Tempo: "
                f"{time.monotonic() - t0:.1f}s"
            )

    return (
        df,
        end - start,
        time.monotonic() - t0
    )


# ============================================================
# GERAR TAREFAS
# ============================================================

def tasks_for(obras):
    rows = []

    for _, r in obras.iterrows():
        for kind in activity(
            r["LISTA"]
        ):
            rows.append({
                **r.to_dict(),
                "ATIVIDADE": kind
            })

    return pd.DataFrame(rows)


# ============================================================
# DATAS PARA PROGRAMACAO
# ============================================================

def schedule_days(
    start,
    mode,
    periods
):
    start = pd.Timestamp(start)

    if mode == "POR DIA":
        end = (
            start
            + pd.Timedelta(
                days=periods * 3 + 14
            )
        )

    elif mode == "POR SEMANA":
        end = (
            start
            + pd.Timedelta(
                weeks=periods
            )
        )

    elif mode == "POR MÊS":
        end = (
            start
            + pd.DateOffset(
                months=periods
            )
            - pd.Timedelta(days=1)
        )

    else:
        end = (
            start
            + pd.Timedelta(days=365)
        )

    days = [
        v.date()
        for v in pd.date_range(
            start,
            end,
            freq="D"
        )
        if v.weekday() < 5
    ]

    if mode == "POR DIA":
        return days[:periods]

    return days


def period_id(date, mode):
    iso = date.isocalendar()

    if mode == "POR SEMANA":
        return (
            f"{iso.year}-S"
            f"{iso.week:02d}"
        )

    if mode == "POR MÊS":
        return date.strftime("%Y-%m")

    return date.isoformat()


# ============================================================
# ALOCACAO COM DISTANCIA RODOVIARIA
# ============================================================

def schedule_tasks(
    tasks,
    routes,
    teams,
    capacities,
    mode,
    periods,
    start,
    period_cap,
    default_daily,
    limit_km,
    route_only
):
    cols = [
        "EQUIPE_PROGRAMADA",
        "ID_EQUIPE",
        "DATA_PROGRAMADA",
        "PERIODO",
        "STATUS_PROGRAMACAO"
    ]

    if tasks.empty:
        return (
            pd.DataFrame(
                columns=list(tasks.columns) + cols
            ),
            pd.DataFrame()
        )

    routes = routes.copy()

    routes["KM_ALOCACAO"] = (
        routes["KM_RODOVIARIO"]
    )

    route_map = defaultdict(list)

    for _, c in routes.iterrows():
        if (
            c["ROTA_STATUS"]
            == "OSRM CONFIRMADA"
            and pd.notna(
                c["KM_RODOVIARIO"]
            )
            and c["KM_RODOVIARIO"] <= limit_km
        ):
            route_map[
                (
                    c["NOTA"],
                    c["ATIVIDADE"]
                )
            ].append(
                c.to_dict()
            )

    capmap = {
        str(r["ID_EQUIPE"]): int(
            r["CAPACIDADE_DIA"]
        )
        for _, r in capacities.iterrows()
    }

    days = schedule_days(
        start,
        mode,
        periods
    )

    used_day = defaultdict(int)
    used_period = defaultdict(int)

    order = {
        "ALTA": 0,
        "MEDIA": 1,
        "BAIXA": 2
    }

    ordered = tasks.assign(
        _ord=tasks[
            "PRIORIDADE_PLANEJAMENTO"
        ].map(order).fillna(3)
    ).sort_values(
        [
            "_ord",
            "MUNICIPIO",
            "NOTA"
        ]
    )

    out = []

    for _, task in ordered.iterrows():
        candidates = route_map.get(
            (
                task["NOTA"],
                task["ATIVIDADE"]
            ),
            []
        )

        options = []

        for c in candidates:
            for day in days:
                period = period_id(
                    day,
                    mode
                )

                key = (
                    c["ID_EQUIPE"],
                    day
                )

                pk = (
                    c["ID_EQUIPE"],
                    period
                )

                cap = int(
                    capmap.get(
                        c["ID_EQUIPE"],
                        default_daily
                    )
                )

                if (
                    used_day[key] < cap
                    and used_period[pk] < period_cap
                ):
                    options.append(
                        (
                            used_period[pk],
                            used_day[key],
                            day,
                            float(
                                c["KM_RODOVIARIO"]
                            ),
                            c,
                            period
                        )
                    )

                    break

        r = task.drop(
            labels=["_ord"]
        ).to_dict()

        if options:
            (
                _,
                _,
                day,
                km,
                c,
                period
            ) = min(
                options,
                key=lambda v: v[:4]
            )

            used_day[
                (
                    c["ID_EQUIPE"],
                    day
                )
            ] += 1

            used_period[
                (
                    c["ID_EQUIPE"],
                    period
                )
            ] += 1

            r.update({
                "EQUIPE_PROGRAMADA": c["EQUIPE"],
                "ID_EQUIPE": c["ID_EQUIPE"],
                "DATA_PROGRAMADA": day.isoformat(),
                "DIA_SEMANA": DAY_NAMES[
                    day.weekday()
                ],
                "PERIODO": period,
                "KM_RODOVIARIO": km,
                "TEMPO_RODOVIARIO_MIN": c[
                    "TEMPO_RODOVIARIO_MIN"
                ],
                "STATUS_PROGRAMACAO": "SUGESTÃO"
            })

        else:
            r.update({
                "EQUIPE_PROGRAMADA": "NÃO ALOCADA",
                "ID_EQUIPE": "",
                "DATA_PROGRAMADA": "",
                "DIA_SEMANA": "",
                "PERIODO": "",
                "KM_RODOVIARIO": np.nan,
                "TEMPO_RODOVIARIO_MIN": np.nan,
                "STATUS_PROGRAMACAO": (
                    "SEM ROTA OU CAPACIDADE"
                )
            })

        out.append(r)

    plan = pd.DataFrame(out)

    load = (
        plan[
            plan[
                "STATUS_PROGRAMACAO"
            ].eq("SUGESTÃO")
        ]
        .groupby(
            [
                "ATIVIDADE",
                "EQUIPE_PROGRAMADA",
                "PERIODO"
            ]
        )
        .size()
        .reset_index(
            name="TAREFAS"
        )
    )

    return plan, load


# ============================================================
# SUPERPONTOS DBSCAN
# ============================================================

def merge_superpoints(
    plan,
    meters
):
    if plan.empty:
        return (
            pd.DataFrame(),
            0
        )

    valid = plan.dropna(
        subset=[
            "LATITUDE",
            "LONGITUDE"
        ]
    ).copy()

    invalid = plan[
        plan["LATITUDE"].isna()
        | plan["LONGITUDE"].isna()
    ].copy()

    if valid.empty:
        return invalid, 0

    grouped = []

    # Agrupamento independente por equipe,
    # data e atividade.
    for _, g in valid.groupby(
        [
            "ID_EQUIPE",
            "DATA_PROGRAMADA",
            "ATIVIDADE"
        ],
        dropna=False,
        sort=False
    ):
        coords = np.radians(
            g[
                [
                    "LATITUDE",
                    "LONGITUDE"
                ]
            ].to_numpy(float)
        )

        labels = DBSCAN(
            eps=meters / 6371000,
            min_samples=1,
            metric="haversine",
            algorithm="ball_tree"
        ).fit(coords).labels_

        clustered = g.assign(
            _cluster=labels
        )

        for _, cluster in clustered.groupby(
            "_cluster",
            sort=False
        ):
            r = cluster.iloc[0].drop(
                labels=["_cluster"]
            ).to_dict()

            original = cluster.drop(
                columns="_cluster"
            ).to_dict("records")

            r["_ORIGINAL_ROWS"] = original

            r["LATITUDE"] = (
                cluster["LATITUDE"].mean()
            )

            r["LONGITUDE"] = (
                cluster["LONGITUDE"].mean()
            )

            r["SUPER_PONTO"] = (
                f"SIM ({len(cluster)} un.)"
                if len(cluster) > 1
                else "NÃO"
            )

            r["NOTAS_AGRUPADAS"] = (
                " | ".join(
                    cluster[
                        "NOTA"
                    ].astype(str)
                )
            )

            grouped.append(r)

    for _, r in invalid.iterrows():
        grouped.append({
            **r.to_dict(),
            "SUPER_PONTO": "NÃO",
            "NOTAS_AGRUPADAS": r["NOTA"],
            "_ORIGINAL_ROWS": [
                r.to_dict()
            ]
        })

    return (
        pd.DataFrame(grouped),
        len(plan) - len(grouped)
    )


# ============================================================
# AUDITORIA DE INCONSISTENCIAS
# ============================================================

def anomalies(
    san,
    lev,
    consolidated,
    teams
):
    issues = []

    for origin, df in [
        ("SANEAMENTO", san),
        ("LEVANTAMENTO", lev)
    ]:
        counts = (
            df.loc[
                df["CHAVE"].ne(""),
                "CHAVE"
            ].value_counts()
        )

        repeated = counts[
            counts > 1
        ]

        for key, count in repeated.items():
            issues.append({
                "NOTA": key,
                "TIPO": (
                    "REPETIDA NA BASE "
                    + origin
                ),
                "DETALHE": f"{count} linhas"
            })

    for _, r in consolidated.iterrows():
        if (
            pd.isna(r["LATITUDE"])
            or pd.isna(r["LONGITUDE"])
        ):
            issues.append({
                "NOTA": r["NOTA"],
                "TIPO": "COORDENADA AUSENTE",
                "DETALHE": r["MUNICIPIO"]
            })

        municipio_san = n(
            r["MUNICIPIO_SANEAMENTO"]
        )

        municipio_lev = n(
            r["MUNICIPIO_LEVANTAMENTO"]
        )

        if (
            municipio_san
            and municipio_lev
            and municipio_san != municipio_lev
        ):
            issues.append({
                "NOTA": r["NOTA"],
                "TIPO": (
                    "DIVERGÊNCIA DE MUNICÍPIO"
                ),
                "DETALHE": (
                    f"{r['MUNICIPIO_SANEAMENTO']} / "
                    f"{r['MUNICIPIO_LEVANTAMENTO']}"
                )
            })

        if (
            r["SAP_FINL"] == "SIM"
            and r["SAP_CANC"] == "SIM"
        ):
            issues.append({
                "NOTA": r["NOTA"],
                "TIPO": (
                    "STATUS SAP CONFLITANTE"
                ),
                "DETALHE": (
                    "FINL e CANC em linhas diferentes"
                )
            })

    return pd.DataFrame(
        issues,
        columns=[
            "NOTA",
            "TIPO",
            "DETALHE"
        ]
    )


# ============================================================
# COMPARACAO HISTORICA
# ============================================================

def compare_history(
    current,
    file
):
    sheets = pd.read_excel(
        io.BytesIO(file.getvalue()),
        sheet_name=None,
        engine="openpyxl",
        dtype=str
    )

    previous = next(
        (
            sheets[x]
            for x in [
                "CONSOLIDADO",
                "AUDITORIA COMPLETA",
                "AUDITORIA"
            ]
            if x in sheets
        ),
        None
    )

    if (
        previous is None
        or "NOTA" not in previous.columns
    ):
        raise ValueError(
            "Histórico deve conter CONSOLIDADO "
            "ou AUDITORIA e coluna NOTA."
        )

    if (
        "PENDENTE_CONTAGEM"
        in previous.columns
    ):
        previous = previous[
            previous[
                "PENDENTE_CONTAGEM"
            ].eq("SIM")
        ]

    old = set(
        previous["NOTA"].map(nota)
    ) - {""}

    new = set(
        current["NOTA"].map(nota)
    ) - {""}

    return (
        pd.DataFrame({
            "NOTA": sorted(
                new - old
            )
        }),
        pd.DataFrame({
            "NOTA": sorted(
                old - new
            )
        })
    )


# ============================================================
# EXCEL
# ============================================================

def safe_excel(df):
    df = df.copy()

    for c in df.columns:
        if df[c].dtype == "object":
            df[c] = df[c].map(
                lambda x: (
                    " | ".join(
                        map(str, x)
                    )
                    if isinstance(x, list)
                    else str(x)
                    if isinstance(x, dict)
                    else x
                )
            )

    return df.drop(
        columns=[
            "_ORIGINAL_ROWS",
            "GEOMETRIA"
        ],
        errors="ignore"
    )


def excel_bytes(sheets):
    out = io.BytesIO()

    with pd.ExcelWriter(
        out,
        engine="openpyxl"
    ) as writer:

        for name, frame in sheets.items():
            df = safe_excel(frame)

            sheet = name[:31]

            df.to_excel(
                writer,
                sheet_name=sheet,
                index=False
            )

            ws = writer.sheets[sheet]

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            ws.sheet_view.showGridLines = False

            for cell in ws[1]:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="002060"
                )

                cell.font = Font(
                    color="FFFFFF",
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center"
                )

            for row in ws.iter_rows(
                min_row=2
            ):
                if (
                    "SUPER_PONTO" in df.columns
                    and str(
                        row[
                            df.columns.get_loc(
                                "SUPER_PONTO"
                            )
                        ].value
                    ).startswith("SIM")
                ):
                    for cell in row:
                        cell.fill = PatternFill(
                            "solid",
                            fgColor="FCE4D6"
                        )

            for col_cells in ws.columns:
                letter = (
                    col_cells[0].column_letter
                )

                lengths = [
                    len(str(c.value or ""))
                    for c in list(
                        col_cells
                    )[:100]
                ]

                ws.column_dimensions[
                    letter
                ].width = min(
                    55,
                    max(
                        13,
                        max(
                            lengths,
                            default=10
                        ) + 2
                    )
                )

    return out.getvalue()


# ============================================================
# POPUP DOS MARCADORES
# ============================================================

def popup(r):
    """Mesmo cartão HTML no mapa Folium e na descrição do KML."""
    def value(key, default="Não informado"):
        x = r.get(key, default)
        try:
            if pd.isna(x):
                return default
        except (ValueError, TypeError):
            pass
        return str(x).strip() or default
    def esc(x):
        return html.escape(str(x), quote=True)
    cat = value("LISTA", "")
    color = COLORS.get(cat, "#173d59")
    status = value("STATUS_LIST", "Pendente")
    nota_id = value("NOTA", "")
    nearest = value("EQUIPE_MAIS_PROXIMA", "Não consultada")
    attributed = value("BASE_ATRIBUIDA", "Sem equipe")
    km = value("DISTANCIA_EQUIPE_KM", "Sem rota OSRM")
    method = value("METODO_DISTANCIA", "OSRM pendente")
    try:
        km = f"{float(km):.2f} km"
    except (TypeError, ValueError):
        pass
    fields = [
        ("Regional", value("REGIONAL")), ("Município", value("MUNICIPIO")),
        ("Levantador", value("EQUIPE_CAMPO", "SEM LEVANTADOR")),
        ("Contrato", value("CONTRATO", "NIP GLOBAL LTDA - EQTL MARANHÃO")),
        ("Status SAP", value("STATUS_SAP")), ("Status LIST", status),
        ("Equipe mais próxima", nearest), ("Equipe atribuída", attributed),
        ("Atividade da equipe", value("ATIVIDADE_EQUIPE")),
        ("Base da equipe", value("CIDADE_BASE")),
        ("Distância rodoviária", km), ("Fonte da distância", method),
        ("Situação", value("STATUS_DISTRIBUICAO")),
        ("Superponto 50 m", value("SUPER_PONTO", "Não informado")),
        ("Notas no agrupamento", value("NOTAS_AGRUPADAS", nota_id)),
        ("Equipes próximas / OSRM", value("RANKING_EQUIPES", "Aguardando consulta")),
    ]
    rows = "".join(
        '<div style="display:grid;grid-template-columns:42% 58%;gap:0;border-bottom:1px solid #dfe5eb;padding:8px 3px">'
        f'<span style="color:#4b5563">{esc(k)}</span><strong style="overflow-wrap:anywhere">{esc(v)}</strong></div>'
        for k,v in fields
    )
    lat,lon = r.get("LATITUDE"),r.get("LONGITUDE")
    links = ""
    try:
        lat,lon=float(lat),float(lon)
        if np.isfinite(lat) and np.isfinite(lon):
            gm=f"https://www.google.com/maps/search/?api=1&query={lat:.7f}%2C{lon:.7f}"
            earth=f"https://earth.google.com/web/search/{lat:.7f},{lon:.7f}"
            links=(f'<a target="_blank" style="display:block;background:#087e90;color:white;padding:11px;text-align:center;margin-top:9px;text-decoration:none;font-weight:bold" href="{gm}">ABRIR IMAGEM DE SATÉLITE ↗</a>'
                   f'<a target="_blank" style="display:block;background:#e5eef2;color:#13607c;padding:9px;text-align:center;margin-top:7px;text-decoration:none" href="{earth}">Abrir no Google Earth Web</a>'
                   f'<a target="_blank" style="display:block;background:#e5eef2;color:#13607c;padding:9px;text-align:center;margin-top:7px;text-decoration:none" href="{gm}">Abrir no Google Maps</a>')
    except (TypeError, ValueError):
        pass
    return (
        '<div style="font:12px Arial,sans-serif;color:#26313c;width:315px;max-height:500px;overflow-y:auto;background:white">'
        f'<div style="background:#173d59;color:white;padding:14px"><small>OBRA / NOTA</small><div style="font-size:21px;font-weight:bold">{esc(nota_id)}</div>'
        f'<small>{esc(cat)}</small></div>'
        f'<div style="padding:10px"><div style="display:inline-block;background:{color};color:white;padding:5px 8px;border-radius:3px;margin-bottom:7px">{esc(status)}</div>'
        f'{rows}{links}</div></div>'
    )


# ============================================================
# GERAR KML
# ============================================================

def kml_bytes(
    works,
    superpoints=None,
    paths=None,
    name="NIP - Obras"
):
    ns = (
        "http://www.opengis.net/kml/2.2"
    )

    ET.register_namespace("", ns)

    def add(parent, tag, value=None):
        x = ET.SubElement(
            parent,
            f"{{{ns}}}{tag}"
        )

        x.text = (
            str(value)
            if value is not None
            else None
        )

        return x

    root = ET.Element(
        f"{{{ns}}}kml"
    )

    doc = add(
        root,
        "Document"
    )

    add(
        doc,
        "name",
        name
    )

    icons = {
        CAT_SAN: "blu-blank.png",
        CAT_LEV: "grn-blank.png",
        CAT_DUP: "purple-blank.png"
    }

    for cat in CATS:
        style = add(
            doc,
            "Style"
        )

        style.set(
            "id",
            "cat" + str(
                CATS.index(cat)
            )
        )

        ic = add(
            add(
                style,
                "IconStyle"
            ),
            "Icon"
        )

        add(
            ic,
            "href",
            (
                "http://maps.google.com/mapfiles/"
                "kml/paddle/"
                + icons[cat]
            )
        )

        folder = add(
            doc,
            "Folder"
        )

        add(
            folder,
            "name",
            cat
        )

        if works.empty:
            continue

        subset = works[
            works[
                "LISTA"
            ].eq(cat)
        ].dropna(
            subset=[
                "LATITUDE",
                "LONGITUDE"
            ]
        )

        for _, r in subset.iterrows():
            pm = add(
                folder,
                "Placemark"
            )

            add(
                pm,
                "name",
                (
                    "SUPERPONTO " + str(r["SUPER_PONTO"])
                    if str(r.get("SUPER_PONTO", "")).startswith("SIM")
                    else "Nota " + str(r["NOTA"])
                )
            )

            add(
                pm,
                "styleUrl",
                "#cat"
                + str(
                    CATS.index(cat)
                )
            )

            add(
                pm,
                "description",
                popup(r)
            )

            add(
                add(
                    pm,
                    "Point"
                ),
                "coordinates",
                (
                    f"{r['LONGITUDE']},"
                    f"{r['LATITUDE']},0"
                )
            )

    # SUPERPONTOS
    if (
        superpoints is not None
        and not superpoints.empty
    ):
        folder = add(
            doc,
            "Folder"
        )

        add(
            folder,
            "name",
            "SUPERPONTOS"
        )

        dados = superpoints[
            superpoints[
                "SUPER_PONTO"
            ].astype(str).str.startswith(
                "SIM"
            )
        ].dropna(
            subset=[
                "LATITUDE",
                "LONGITUDE"
            ]
        )

        for _, r in dados.iterrows():
            pm = add(
                folder,
                "Placemark"
            )

            add(
                pm,
                "name",
                (
                    "SUPERPONTO "
                    + str(
                        r["SUPER_PONTO"]
                    )
                )
            )

            add(
                pm,
                "styleUrl",
                "#cat"
                + str(
                    CATS.index(
                        r["LISTA"]
                    )
                )
            )

            add(
                pm,
                "description",
                popup(r)
            )

            add(
                add(
                    pm,
                    "Point"
                ),
                "coordinates",
                (
                    f"{r['LONGITUDE']},"
                    f"{r['LATITUDE']},0"
                )
            )

    # TRAJETOS OSRM CONFIRMADOS
    if (
        paths is not None
        and not paths.empty
    ):
        folder = add(
            doc,
            "Folder"
        )

        add(
            folder,
            "name",
            "TRAJETOS OSRM CONFIRMADOS"
        )

        for _, r in paths.iterrows():
            coords = r.get(
                "GEOMETRIA",
                []
            )

            if (
                not isinstance(coords, list)
                or len(coords) < 2
            ):
                continue

            pm = add(
                folder,
                "Placemark"
            )

            add(
                pm,
                "name",
                (
                    f"{r['EQUIPE']} "
                    f"| {r['NOTA']}"
                )
            )

            line = add(
                pm,
                "LineString"
            )

            add(
                line,
                "tessellate",
                "1"
            )

            add(
                line,
                "coordinates",
                " ".join(
                    f"{x},{y},0"
                    for x, y in coords
                )
            )

    return ET.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True
    )


# ============================================================
# LISTA CONTINUA - EXPORTACAO POR EQUIPE MAIS PROXIMA
# ============================================================

def classificar_tipo_obra(categoria):
    """Classificacao original da nota, nunca altera sua contagem."""
    return {
        CAT_SAN: "SANEAMENTO",
        CAT_LEV: "LEVANTAMENTO",
        CAT_DUP: "SANEAMENTO E LEVANTAMENTO",
    }.get(categoria, str(categoria))


def escolher_equipes_proximas(obras, equipes):
    """Atribui UMA nota a UMA equipe mais proxima elegivel.

    A referencia e Haversine, nao distancia rodoviaria. Para obras
    de ambas as atividades, compara equipes dos dois grupos e preserva
    a classificacao original da nota no Excel e no KML.
    """
    saida = obras.copy().reset_index(drop=True)
    for campo in ["TIPO_OBRA", "BASE_ATRIBUIDA", "ID_EQUIPE", "ATIVIDADE_EQUIPE", "CIDADE_BASE", "METODO_DISTANCIA"]:
        saida[campo] = ""
    saida["DISTANCIA_EQUIPE_KM"] = np.nan
    saida["TIPO_OBRA"] = saida["LISTA"].map(classificar_tipo_obra)
    saida["STATUS_DISTRIBUICAO"] = "SEM COORDENADAS OU EQUIPE"
    for tipo in ["SANEAMENTO", "LEVANTAMENTO"]:
        equipe_tipo = equipes.loc[
            equipes["TIPO_EQUIPE"].eq(tipo),
            ["ID_EQUIPE", "EQUIPE", "CIDADE_BASE", "LAT_EQUIPE", "LON_EQUIPE"]
        ].dropna(subset=["LAT_EQUIPE", "LON_EQUIPE"]).reset_index(drop=True)
        if equipe_tipo.empty:
            continue
        mascara = (
            saida["LISTA"].map(lambda cat: tipo in activity(cat))
            & saida["LATITUDE"].notna()
            & saida["LONGITUDE"].notna()
        )
        indices = saida.index[mascara].to_numpy()
        if not len(indices):
            continue
        arvore = BallTree(
            np.radians(equipe_tipo[["LAT_EQUIPE", "LON_EQUIPE"]].to_numpy(float)),
            metric="haversine",
        )
        distancias, posicoes = arvore.query(
            np.radians(saida.loc[indices, ["LATITUDE", "LONGITUDE"]].to_numpy(float)),
            k=1,
        )
        for i, indice in enumerate(indices):
            km = float(distancias[i, 0] * R)
            anterior = saida.at[indice, "DISTANCIA_EQUIPE_KM"]
            if pd.isna(anterior) or km < float(anterior):
                equipe = equipe_tipo.iloc[int(posicoes[i, 0])]
                saida.at[indice, "DISTANCIA_EQUIPE_KM"] = round(km, 2)
                saida.at[indice, "BASE_ATRIBUIDA"] = equipe["EQUIPE"]
                saida.at[indice, "ID_EQUIPE"] = equipe["ID_EQUIPE"]
                saida.at[indice, "ATIVIDADE_EQUIPE"] = tipo
                saida.at[indice, "CIDADE_BASE"] = equipe["CIDADE_BASE"]
                saida.at[indice, "METODO_DISTANCIA"] = "LINHA RETA (HAVERSINE)"
                saida.at[indice, "STATUS_DISTRIBUICAO"] = "EQUIPE MAIS PROXIMA"
    saida["EQUIPE_PROGRAMADA"] = saida["BASE_ATRIBUIDA"]
    return saida


def agrupar_rota_equipe(df, metros):
    """DBSCAN por equipe + classificacao: uma linha por superponto KML;
    Excel mantem uma linha para cada nota e marca superponto e ordem.
    Segue a organizacao da Lista Continua sem alterar os modulos comuns.
    """
    if df.empty:
        return df.copy(), df.copy()
    kml_rows, excel_rows = [], []
    for tipo_obra, subset in df.groupby("TIPO_OBRA", sort=False):
        rad = np.radians(subset[["LATITUDE", "LONGITUDE"]].to_numpy(float))
        labels = DBSCAN(
            eps=float(metros) / 6371000.0,
            min_samples=1,
            metric="haversine",
            algorithm="ball_tree",
        ).fit_predict(rad)
        partes = []
        for _, cluster in subset.assign(_grupo=labels).groupby("_grupo", sort=False):
            primeira = cluster.iloc[0].drop(labels="_grupo").to_dict()
            primeira["LATITUDE"] = float(cluster["LATITUDE"].mean())
            primeira["LONGITUDE"] = float(cluster["LONGITUDE"].mean())
            primeira["SUPER_PONTO"] = f"SIM ({len(cluster)} Obras)" if len(cluster) > 1 else "NÃO"
            primeira["NOTAS_AGRUPADAS"] = " | ".join(cluster["NOTA"].astype(str))
            primeira["_ORIGINAL_ROWS"] = cluster.drop(columns="_grupo").to_dict("records")
            partes.append(primeira)
        # Uma rota heuristica de vizinho proximo por equipe, sem alegar OSRM.
        restante = partes[:]
        atual = None
        ordem = 0
        while restante:
            if atual is None:
                pos = min(range(len(restante)), key=lambda j: (str(restante[j]["MUNICIPIO"]), str(restante[j]["NOTA"])))
                km_anterior = 0.0
            else:
                ds = [float(haversine(atual["LATITUDE"], atual["LONGITUDE"], x["LATITUDE"], x["LONGITUDE"])) for x in restante]
                pos = int(np.argmin(ds))
                km_anterior = round(ds[pos], 2)
            item = restante.pop(pos)
            ordem += 1
            item["ORDEM"] = ordem
            item["DISTANCIA_PONTO_ANTERIOR_KM"] = km_anterior
            kml_rows.append(item)
            for original in item["_ORIGINAL_ROWS"]:
                linha = dict(original)
                linha["ORDEM"] = ordem
                linha["SUPER_PONTO"] = item["SUPER_PONTO"]
                linha["NOTAS_AGRUPADAS"] = item["NOTAS_AGRUPADAS"]
                linha["DISTANCIA_PONTO_ANTERIOR_KM"] = km_anterior
                linha["DIA_SEMANA"] = "NÃO PROGRAMADA"
                linha["DIA_MES"] = ""
                excel_rows.append(linha)
            atual = item
    return pd.DataFrame(excel_rows), pd.DataFrame(kml_rows)


def nome_seguro_equipe(equipe, identificador):
    base = n(equipe).replace(" ", "_")
    base = re.sub(r"[^A-Z0-9_-]", "", base)[:54].strip("_") or "EQUIPE"
    codigo = hashlib.sha1(str(identificador).encode("utf-8")).hexdigest()[:7]
    return f"{base}_{codigo}"


def gerar_zips_lista_continua_nip(obras, equipes, metros, kml_geral, progresso=None):
    """Retorna ZIP Excel/KML + resumo usando TODO o backlog pendente."""
    distribuicao = escolher_equipes_proximas(obras, equipes)
    com_eq = distribuicao.loc[
        distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")
    ].copy()
    sem_eq = distribuicao.loc[
        ~distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")
    ].copy()
    resumo = (
        com_eq.groupby(["ATIVIDADE_EQUIPE", "BASE_ATRIBUIDA", "ID_EQUIPE"], dropna=False)
        .agg(OBRAS=("NOTA", "nunique"), MUNICIPIOS=("MUNICIPIO", "nunique"))
        .reset_index()
    )
    data_fmt = datetime.now().strftime("%d.%m.%Y")
    total_equipes = com_eq["ID_EQUIPE"].nunique()
    if progresso is not None:
        progresso(0, total_equipes, "Iniciando arquivos gerais")
    excel_zip = io.BytesIO()
    kml_zip = io.BytesIO()
    with zipfile.ZipFile(excel_zip, "w", zipfile.ZIP_DEFLATED) as zx, zipfile.ZipFile(kml_zip, "w", zipfile.ZIP_DEFLATED) as zk:
        zx.writestr(f"Resumo_Operacional - {data_fmt}.xlsx", excel_bytes({"Resumo Operacional": resumo, "Sem Equipe": sem_eq}))
        zx.writestr(f"Demanda_ListaContinua_Total - {data_fmt}.xlsx", excel_bytes({"Obras Roteirizadas": distribuicao}))
        zk.writestr(f"ROTA_TOTAL - {data_fmt}.kml", kml_geral)
        for numero_eq, (ident, grupo) in enumerate(com_eq.groupby("ID_EQUIPE", sort=True), start=1):
            info = grupo.iloc[0]
            atividade = info["ATIVIDADE_EQUIPE"]
            safe = nome_seguro_equipe(info["BASE_ATRIBUIDA"], ident)
            excel_rows, kml_rows = agrupar_rota_equipe(grupo, metros)
            # Uma linha por nota original, inclusive as agrupadas em superpontos.
            primeiras = ["BASE_ATRIBUIDA", "TIPO_OBRA", "ATIVIDADE_EQUIPE", "DIA_SEMANA", "DIA_MES", "SUPER_PONTO", "ORDEM", "DISTANCIA_PONTO_ANTERIOR_KM", "NOTA", "MUNICIPIO", "REGIONAL", "DISTANCIA_EQUIPE_KM", "METODO_DISTANCIA"]
            outras = [c for c in excel_rows.columns if c not in primeiras and not c.startswith("_")]
            excel_rows = excel_rows[[c for c in primeiras if c in excel_rows.columns] + outras]
            zx.writestr(f"Rotas_{data_fmt}/{atividade}/Rota_{safe}.xlsx", excel_bytes({"Obras Roteirizadas": excel_rows}))
            # O KML da equipe tem um marcador por superponto ou obra individual.
            zk.writestr(f"KML_{data_fmt}/{atividade}/Rota_{safe}.kml", kml_bytes(kml_rows, name=f"Rota {info['BASE_ATRIBUIDA']}"))
            if progresso is not None:
                progresso(numero_eq, total_equipes, str(info["BASE_ATRIBUIDA"]))
    return excel_zip.getvalue(), kml_zip.getvalue(), distribuicao, resumo, sem_eq


# ============================================================
# CONTROLE E CONFERENCIA DA DISTRIBUICAO
# ============================================================

@st.cache_data(show_spinner=False, max_entries=8)
def distribuicao_lista_cache(obras, equipes):
    """Somente memoriza a funcao original de equipe mais proxima."""
    return escolher_equipes_proximas(obras, equipes)


def conferir_distribuicao(obras, distribuicao):
    """Nao altera nenhuma nota: so confere a particao e a integridade."""
    originais = obras["NOTA"].astype(str).tolist()
    atribuidas = distribuicao["NOTA"].astype(str).tolist()
    com = distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")
    qtd_origem = len(set(originais))
    qtd_dist = int(com.sum())
    qtd_sem = int((~com).sum())
    duplicadas = int(pd.Series(atribuidas).duplicated().sum())
    faltantes = sorted(set(originais) - set(atribuidas))
    extras = sorted(set(atribuidas) - set(originais))
    correto = (qtd_origem == qtd_dist + qtd_sem and duplicadas == 0
               and not faltantes and not extras and len(originais) == qtd_origem)
    return {"OK": correto, "PENDENTES": qtd_origem, "DISTRIBUIDAS": qtd_dist,
            "SEM_EQUIPE": qtd_sem, "DUPLICADAS_INDEVIDAS": duplicadas,
            "FALTANTES": faltantes, "EXTRAS": extras}


def resumo_detalhado_equipes(distribuicao, equipes):
    """Um registro por equipe, inclusive quando tem zero notas."""
    resumo = equipes[["ID_EQUIPE", "EQUIPE", "TIPO_EQUIPE", "CIDADE_BASE"]].copy()
    atrib = distribuicao[distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")].copy()
    if not atrib.empty:
        agrupado = atrib.groupby("ID_EQUIPE").agg(
            OBRAS=("NOTA", "nunique"), MUNICIPIOS=("MUNICIPIO", "nunique"),
            DISTANCIA_MEDIA_LINHA_RETA_KM=("DISTANCIA_EQUIPE_KM", "mean"),
            DISTANCIA_MAX_LINHA_RETA_KM=("DISTANCIA_EQUIPE_KM", "max")
        ).reset_index()
        tipos = atrib.groupby(["ID_EQUIPE", "TIPO_OBRA"]).size().unstack(fill_value=0)
        tipos = tipos.rename(columns={
            "SANEAMENTO": "OBRAS_SANEAMENTO",
            "LEVANTAMENTO": "OBRAS_LEVANTAMENTO",
            "SANEAMENTO E LEVANTAMENTO": "OBRAS_AMBAS"
        }).reset_index()
        resumo = resumo.merge(agrupado, on="ID_EQUIPE", how="left").merge(tipos, on="ID_EQUIPE", how="left")
    for nome in ["OBRAS", "MUNICIPIOS", "OBRAS_SANEAMENTO", "OBRAS_LEVANTAMENTO", "OBRAS_AMBAS"]:
        if nome not in resumo.columns:
            resumo[nome] = 0
        resumo[nome] = resumo[nome].fillna(0).astype(int)
    for nome in ["DISTANCIA_MEDIA_LINHA_RETA_KM", "DISTANCIA_MAX_LINHA_RETA_KM"]:
        if nome not in resumo.columns:
            resumo[nome] = np.nan
        resumo[nome] = resumo[nome].round(1)
    total = max(int(resumo["OBRAS"].sum()), 1)
    resumo["PARTICIPACAO_PERCENTUAL"] = (100 * resumo["OBRAS"] / total).round(1)
    return resumo.sort_values(["TIPO_EQUIPE", "OBRAS"], ascending=[True, False]).reset_index(drop=True)


def tabela_sem_equipe(distribuicao):
    sem = distribuicao[~distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")].copy()
    if sem.empty:
        sem["MOTIVO_SEM_EQUIPE"] = pd.Series(dtype="str")
        return sem
    coords = sem["LATITUDE"].notna() & sem["LONGITUDE"].notna()
    sem["MOTIVO_SEM_EQUIPE"] = np.where(coords, "SEM EQUIPE ELEGIVEL COM LOCALIZACAO", "SEM COORDENADAS VALIDAS")
    return sem


def parametros_analise(registro):
    return pd.DataFrame([{"PARAMETRO": k, "VALOR": str(v)} for k, v in registro.items()])


def inventario_zips(distribuicao, equipes, data_fmt):
    com = distribuicao[distribuicao["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")]
    rows = [
        {"PACOTE":"EXCEL", "ARQUIVO":f"Resumo_Operacional - {data_fmt}.xlsx", "NOTAS":len(distribuicao)},
        {"PACOTE":"EXCEL", "ARQUIVO":f"Demanda_ListaContinua_Total - {data_fmt}.xlsx", "NOTAS":len(distribuicao)},
        {"PACOTE":"KML", "ARQUIVO":f"ROTA_TOTAL - {data_fmt}.kml", "NOTAS":len(distribuicao)}
    ]
    for ident, grupo in com.groupby("ID_EQUIPE", sort=True):
        primeira = grupo.iloc[0]
        safe = nome_seguro_equipe(primeira["BASE_ATRIBUIDA"], ident)
        caminho = f"{primeira['ATIVIDADE_EQUIPE']}"
        rows.append({"PACOTE":"EXCEL", "ARQUIVO":f"Rotas_{data_fmt}/{caminho}/Rota_{safe}.xlsx", "NOTAS":len(grupo)})
        rows.append({"PACOTE":"KML", "ARQUIVO":f"KML_{data_fmt}/{caminho}/Rota_{safe}.kml", "NOTAS":len(grupo)})
    return pd.DataFrame(rows)


@st.cache_data(show_spinner=False, max_entries=3)
def gerar_zips_lista_cache(obras, equipes, metros, kml_geral, _indicador=None):
    """Mesmo gerador original; cache pela entrada e raio dos superpontos."""
    return gerar_zips_lista_continua_nip(obras, equipes, metros, kml_geral, progresso=_indicador)


def montar_zips_com_progresso(obras, equipes, metros, kml_geral):
    inicio = time.monotonic()
    barra = st.progress(0, text="Iniciando exportação por equipe...")
    relogio = st.empty()

    def progresso(concluidas, total, equipe):
        decorrido = time.monotonic() - inicio
        percentual = concluidas / total if total else 1.0
        restante = decorrido / concluidas * (total - concluidas) if concluidas else None
        barra.progress(min(1.0, percentual), text=f"Equipe: {equipe} — {concluidas}/{total}")
        relogio.caption(
            f"⏱️ Decorrido: {tempo_legivel(decorrido)} | "
            f"⌛ Restante estimado: {tempo_legivel(restante)}"
        )

    arquivos = gerar_zips_lista_cache(obras, equipes, metros, kml_geral, _indicador=progresso)
    decorrido = time.monotonic() - inicio
    barra.progress(1.0, text="ZIPs concluídos (100%)")
    relogio.caption(f"⏱️ Tempo total: {tempo_legivel(decorrido)} | Restante: 00:00")
    return arquivos



# ============================================================
# MOTOR DE DISTRIBUICAO: OSRM + GRUPOS INSEPARAVEIS DE 50 M
# ============================================================

def preparar_grupos_50m(obras, metros=50):
    """Agrupa pontos de todas as categorias ANTES de escolher a equipe.

    Um cluster com atividades incompatíveis só pode ser distribuído
    integralmente se houver uma equipe apta a ambas (quando existir).
    """
    df = obras.drop_duplicates("NOTA").copy().reset_index(drop=True)
    df["GRUPO_50M"] = ""
    ok = df["LATITUDE"].notna() & df["LONGITUDE"].notna()
    if ok.any():
        coords = np.radians(df.loc[ok, ["LATITUDE", "LONGITUDE"]].to_numpy(float))
        labels = DBSCAN(eps=float(metros) / 6371008.8, min_samples=1,
                        metric="haversine", algorithm="ball_tree").fit_predict(coords)
        df.loc[ok, "GRUPO_50M"] = [f"G{int(i)+1:06d}" for i in labels]
    df.loc[~ok, "GRUPO_50M"] = [f"SEM_COORD_{i}" for i in df.index[~ok]]
    return df


def candidatos_grupos(grupos, equipes, quantidade=3):
    """Pré-seleção geográfica. Nunca confundir esta métrica com distância OSRM."""
    result = []
    valid_teams = equipes.dropna(subset=["LAT_EQUIPE", "LON_EQUIPE"])
    for group_id, sub in grupos.groupby("GRUPO_50M", sort=False):
        if sub["LATITUDE"].isna().any() or sub["LONGITUDE"].isna().any():
            continue
        tipos = [set(activity(cat)) for cat in sub["LISTA"]]
        admissible = set.intersection(*tipos) if tipos else set()
        elig = valid_teams[valid_teams["TIPO_EQUIPE"].isin(admissible)].copy()
        if elig.empty:
            continue
        lat = float(sub["LATITUDE"].mean())
        lon = float(sub["LONGITUDE"].mean())
        elig["KM_RETA_PREFILTRO"] = haversine(lat, lon, elig["LAT_EQUIPE"].to_numpy(float), elig["LON_EQUIPE"].to_numpy(float))
        for _, e in elig.nsmallest(int(quantidade), "KM_RETA_PREFILTRO").iterrows():
            result.append({"GRUPO_50M":group_id,"ID_EQUIPE":e["ID_EQUIPE"],
                           "EQUIPE":e["EQUIPE"],"TIPO_EQUIPE":e["TIPO_EQUIPE"],
                           "CIDADE_BASE":e["CIDADE_BASE"],"LAT_OBRA":lat,"LON_OBRA":lon,
                           "LAT_EQUIPE":float(e["LAT_EQUIPE"]),"LON_EQUIPE":float(e["LON_EQUIPE"]),
                           "KM_RETA_PREFILTRO":float(e["KM_RETA_PREFILTRO"]),
                           "NOTAS_GRUPO":len(sub),"KM_OSRM":np.nan,"MIN_OSRM":np.nan,
                           "STATUS_OSRM":"PENDENTE"})
    return pd.DataFrame(result, columns=["GRUPO_50M","ID_EQUIPE","EQUIPE","TIPO_EQUIPE","CIDADE_BASE",
                                         "LAT_OBRA","LON_OBRA","LAT_EQUIPE","LON_EQUIPE",
                                         "KM_RETA_PREFILTRO","NOTAS_GRUPO","KM_OSRM","MIN_OSRM","STATUS_OSRM"])


def atribuir_por_osrm(grupos, candidatos, equipes, limite=999, balancear=True, limite_km=300):
    """Alocação atômica por cluster, exclusivamente com OSRM confirmado.

    Limites configuraveis de obras e km por equipe. O consumo em km e uma
    estimativa conservadora de ida e volta BASE-GRUPO-BASE (2 x OSRM).
    Nao e quilometragem de uma rota sequencial otimizada.
    Nenhuma distancia Haversine e usada como km rodoviario.
    """
    out = grupos.copy()
    for col_name in ["BASE_ATRIBUIDA","ID_EQUIPE","ATIVIDADE_EQUIPE","CIDADE_BASE",
                     "EQUIPE_MAIS_PROXIMA","RANKING_EQUIPES","METODO_DISTANCIA"]:
        out[col_name] = ""
    out["DISTANCIA_EQUIPE_KM"] = np.nan
    out["KM_IDA_VOLTA_GRUPO"] = np.nan
    out["KM_TOTAL_ESTIMADO_EQUIPE"] = np.nan
    out["STATUS_DISTRIBUICAO"] = "SEM ROTA OSRM CONFIRMADA"
    out["TIPO_OBRA"] = out["LISTA"].map(classificar_tipo_obra)
    out["SUPER_PONTO"] = out["GRUPO_50M"]
    out["NOTAS_AGRUPADAS"] = out.groupby("GRUPO_50M")["NOTA"].transform(lambda x: ", ".join(x.astype(str)))
    load = {str(k):0 for k in equipes["ID_EQUIPE"]}
    km_load = {str(k):0.0 for k in equipes["ID_EQUIPE"]}
    if candidatos.empty:
        return out
    valid = candidatos[(candidatos["STATUS_OSRM"] == "CONFIRMADA") & candidatos["KM_OSRM"].notna()].copy()
    # Grupos com menos opções primeiro; grupos grandes antes de pequenos.
    count = valid.groupby("GRUPO_50M").size().to_dict()
    order = sorted(out["GRUPO_50M"].unique(), key=lambda g:(count.get(g,0), -int((out["GRUPO_50M"]==g).sum()),g))
    for gid in order:
        mask = out["GRUPO_50M"].eq(gid)
        size = int(mask.sum())
        options = valid[valid["GRUPO_50M"].eq(gid)].sort_values("KM_OSRM")
        if options.empty:
            out.loc[mask,"STATUS_DISTRIBUICAO"] = "SEM ROTA OSRM CONFIRMADA"
            continue
        ranking = " | ".join(f"{r.EQUIPE}: {r.KM_OSRM:.2f} km" for r in options.itertuples())
        out.loc[mask,"RANKING_EQUIPES"] = ranking
        out.loc[mask,"EQUIPE_MAIS_PROXIMA"] = str(options.iloc[0]["EQUIPE"])
        capacidade_ok = options[options["ID_EQUIPE"].map(
            lambda i: load[str(i)] + size <= int(limite)
        )]
        if capacidade_ok.empty:
            out.loc[mask,"STATUS_DISTRIBUICAO"] = "LIMITE DE OBRAS / GRUPO INSEPARÁVEL"
            continue
        # Cada grupo de ate 50 m constitui uma visita. Para nao subestimar
        # deslocamento quando falta a rota sequencial, reservar ida e volta
        # entre a base da equipe e o centro do grupo, com distancia OSRM.
        eligible = capacidade_ok[capacidade_ok.apply(
            lambda r: km_load[str(r["ID_EQUIPE"])] + 2.0 * float(r["KM_OSRM"])
            <= float(limite_km) + 1e-9, axis=1
        )]
        if eligible.empty:
            out.loc[mask,"STATUS_DISTRIBUICAO"] = "LIMITE DE KM ESTIMADOS DA EQUIPE"
            continue
        if balancear:
            minimum = float(options.iloc[0]["KM_OSRM"])
            # Não deslocar a equipe para um destino muito distante só para equilibrar.
            competitive = eligible[eligible["KM_OSRM"] <= max(minimum*1.25, minimum+2.0)]
            if not competitive.empty:
                eligible = competitive
            chosen = min((r for _,r in eligible.iterrows()),
                         key=lambda r:(load[str(r["ID_EQUIPE"])],float(r["KM_OSRM"])))
        else:
            chosen = eligible.iloc[0]
        eq_id = str(chosen["ID_EQUIPE"])
        load[eq_id] += size
        km_visita = 2.0 * float(chosen["KM_OSRM"])
        km_load[eq_id] += km_visita
        out.loc[mask,"KM_IDA_VOLTA_GRUPO"] = round(km_visita, 2)
        out.loc[mask,"BASE_ATRIBUIDA"] = chosen["EQUIPE"]
        out.loc[mask,"ID_EQUIPE"] = eq_id
        out.loc[mask,"ATIVIDADE_EQUIPE"] = chosen["TIPO_EQUIPE"]
        out.loc[mask,"CIDADE_BASE"] = chosen["CIDADE_BASE"]
        out.loc[mask,"DISTANCIA_EQUIPE_KM"] = float(chosen["KM_OSRM"])
        out.loc[mask,"METODO_DISTANCIA"] = "OSRM RODOVIÁRIO CONFIRMADO"
        out.loc[mask,"STATUS_DISTRIBUICAO"] = "EQUIPE MAIS PROXIMA"  # Compatibilidade com exports originais
    # Repetir o total por equipe nos registros apenas para consulta/exportacao.
    # Nao somar esta coluna por NOTA, pois um grupo pode conter varias notas.
    out["KM_TOTAL_ESTIMADO_EQUIPE"] = out["ID_EQUIPE"].map(km_load)
    out.loc[out["ID_EQUIPE"].eq(""), "KM_TOTAL_ESTIMADO_EQUIPE"] = np.nan
    out["EQUIPE_PROGRAMADA"] = out["BASE_ATRIBUIDA"]
    return out


def validar_distribuicao_nova(pendentes, distribuicao, limite, limite_km=300):
    base = set(pendentes["NOTA"].astype(str))
    dest = distribuicao["NOTA"].astype(str)
    alocados = distribuicao[distribuicao["STATUS_DISTRIBUICAO"]=="EQUIPE MAIS PROXIMA"]
    cargas = alocados.groupby("ID_EQUIPE")["NOTA"].nunique()
    split = alocados.groupby("GRUPO_50M")["ID_EQUIPE"].nunique()
    # Um superponto e uma visita: nao somar os km para cada nota agrupada.
    visitas = alocados.drop_duplicates(["ID_EQUIPE", "GRUPO_50M"])
    km_equipes = visitas.groupby("ID_EQUIPE")["KM_IDA_VOLTA_GRUPO"].sum()
    return {"Notas originais":len(base),"Notas únicas no resultado":len(set(dest)),
            "Duplicações indevidas":int(dest.duplicated().sum()),
            "Notas ausentes":len(base-set(dest)),
            "Equipes acima do limite":int((cargas>limite).sum()),
            "Equipes acima do limite KM":int((km_equipes>float(limite_km)+1e-6).sum()),
            "KM estimados reservados":round(float(km_equipes.sum()),2),
            "Grupos 50 m divididos":int((split>1).sum()),
            "Alocadas com OSRM":len(alocados),"Sem alocação":len(distribuicao)-len(alocados)}


def zip_distribuicao_osrm(df, equipes, metros, kml_geral, progresso=None):
    """ZIPs são produzidos da atribuição OSRM já validada; nunca recalcular por Haversine."""
    distrib = df.copy()
    com = distrib[distrib["STATUS_DISTRIBUICAO"]=="EQUIPE MAIS PROXIMA"].copy()
    sem = distrib[distrib["STATUS_DISTRIBUICAO"]!="EQUIPE MAIS PROXIMA"].copy()
    resumo = resumo_detalhado_equipes(distrib,equipes)
    stamp = datetime.now().strftime("%d.%m.%Y")
    ex,km=io.BytesIO(),io.BytesIO()
    total=int(com["ID_EQUIPE"].nunique())
    with zipfile.ZipFile(ex,"w",zipfile.ZIP_DEFLATED) as zx,zipfile.ZipFile(km,"w",zipfile.ZIP_DEFLATED) as zk:
        zx.writestr(f"Resumo_Operacional - {stamp}.xlsx",excel_bytes({"Resumo":resumo,"Sem equipe":sem}))
        zx.writestr(f"Demanda_ListaContinua_Total - {stamp}.xlsx",excel_bytes({"Distribuicao":distrib}))
        zk.writestr(f"ROTA_TOTAL - {stamp}.kml",kml_geral)
        for idx,(eqid,group) in enumerate(com.groupby("ID_EQUIPE",sort=True),start=1):
            r=group.iloc[0]
            tipo=str(r["ATIVIDADE_EQUIPE"])
            filename=nome_seguro_equipe(r["BASE_ATRIBUIDA"],eqid)
            excel_rows,kml_rows=agrupar_rota_equipe(group,metros)
            zx.writestr(f"Rotas_{stamp}/{tipo}/Rota_{filename}.xlsx",excel_bytes({"Obras Roteirizadas":excel_rows}))
            zk.writestr(f"KML_{stamp}/{tipo}/Rota_{filename}.kml",kml_bytes(kml_rows,name=f"Equipe {r['BASE_ATRIBUIDA']}"))
            if progresso: progresso(idx,total,str(r["BASE_ATRIBUIDA"]))
    return ex.getvalue(),km.getvalue()


# ============================================================
# NOVA INTERFACE
# ============================================================

st.title("📍 NIP | Análise Cruzada e Distribuição Rodoviária")
st.caption("Saneamento + Levantamento + equipes | OSRM real | 50 m inseparáveis | distribuição controlada")

def zerar_tudo():
    for k in list(st.session_state):
        if k.startswith("nipv2_"):
            del st.session_state[k]
    st.session_state["nipv2_versao"] = st.session_state.get("nipv2_versao",0)+1

versao = st.session_state.get("nipv2_versao",0)
st.button("🗑️ Zerar análise",on_click=zerar_tudo,help="Limpa arquivos, rotas OSRM e distribuição desta página")
with st.expander("Regras de inclusão e atribuição"):
    st.markdown("""As bases são cruzadas por **NOTA = PROTOCOLO**, contando a nota uma única vez.
    FINL e CANC são excluídas; permanecem os demais filtros originais de LIST, contrato,
    colunas M/N e obras já em campo (B/C). Obras a até **50 m** formam grupos inseparáveis.
    Distâncias exibidas como rodoviárias só aparecem após confirmação OSRM.
    A distribuição equilibra a carga entre equipes com rotas competitivas, respeitando os limites de obras e quilometragem selecionados.
    **Não há garantia matemática de trajetos sem cruzamentos** sem otimização global das rotas;
    esta versão minimiza deslocamentos pela proximidade rodoviária e agrupamento territorial.""")

up1,up2,up3=st.columns(3)
with up1: san_file=st.file_uploader("Base SANEAMENTO",type=["xlsx","csv"],key=f"nipv2_san_{versao}")
with up2: lev_file=st.file_uploader("Base LEVANTAMENTO",type=["xlsx"],key=f"nipv2_lev_{versao}")
with up3: team_file=st.file_uploader("Localização das EQUIPES",type=["xlsx"],key=f"nipv2_team_{versao}")
if not all((san_file,lev_file,team_file)):
    st.info("Envie as três bases para iniciar a análise.")
    st.stop()

signature=hashlib.sha256(san_file.getvalue()+lev_file.getvalue()+team_file.getvalue()).hexdigest()
if st.session_state.get("nipv2_signature")!=signature:
    for k in ["nipv2_data","nipv2_candidates","nipv2_zip","nipv2_signature"]:
        st.session_state.pop(k,None)
    st.session_state["nipv2_signature"]=signature

if st.button("🚀 Cruzar e validar as três bases",type="primary"):
    with st.spinner("Lendo arquivos, aplicando exclusões e cruzando as notas..."):
        try:
            san,san_name=load_san(san_file)
            lev,lev_info=load_lev(lev_file)
            teams=load_teams(team_file)
            audit=consolidate(san,lev)
            issues=anomalies(san,lev,audit,teams)
            st.session_state["nipv2_data"]={"audit":audit,"teams":teams,"issues":issues,"san_name":san_name,"lev_info":lev_info}
        except Exception as exc:
            st.error(f"Falha ao processar as bases: {exc}")
if "nipv2_data" not in st.session_state:
    st.stop()

data=st.session_state["nipv2_data"]
audit,teams=data["audit"],data["teams"]
pending=audit[audit["PENDENTE_CONTAGEM"].eq("SIM")].copy().reset_index(drop=True)
finl=audit[audit["SAP_FINL"].eq("SIM")]
canc=audit[audit["SAP_CANC"].eq("SIM")]
field=audit[audit["JA_EM_CAMPO"].eq("SIM")]
duplicadas=pending[pending["DUPLICADA"].eq("SIM")]

cols=st.columns(6)
for c,(title,val) in zip(cols,[
    ("Saneamento",int((pending["LISTA"]==CAT_SAN).sum())),
    ("Levantamento",int((pending["LISTA"]==CAT_LEV).sum())),
    ("Nas duas bases",len(duplicadas)),("Pendentes únicas",len(pending)),
    ("SAP FINL",len(finl)),("SAP CANC",len(canc))]):
    c.metric(title,f"{val:,}".replace(",","."))

st.subheader("1. Agrupamento e cálculo de distância OSRM")
a,b=st.columns(2)
with a:
    limit=int(st.number_input(
        "Quantidade máxima de obras por equipe",
        min_value=1,max_value=999,value=999,step=1,
        help="Selecione de 1 a 999 notas por equipe. Notas a ate 50 m permanecem juntas."
    ))
with b:
    distancia_extra=st.checkbox(
        "Permitir distância extra (acima de 300 km)",value=False,
        help="Desbloqueia valores acima do limite padrao de 300 km por equipe."
    )
    limite_km=int(st.number_input(
        "Quilometragem máxima estimada por equipe (km)",
        min_value=1,max_value=10000 if distancia_extra else 300,
        value=500 if distancia_extra else 300,step=1,
        help="Reserva 2 x distancia OSRM base-grupo para cada grupo atendido. "
             "Nao equivale ao total de uma rota sequencial otimizada."
    ))
st.caption(
    f"Limites selecionados: **{limit} obras/equipe** e **{limite_km} km/equipe**. "
    "Os km representam ida e volta rodoviaria estimada entre a base e cada grupo "
    "de obras (ate 50 m), e nao um roteiro sequencial definitivo."
)
a,b=st.columns(2)
with a: n_near=st.slider("Equipes candidatas por grupo",1,8,3)
with b: batch=st.slider("Consultas OSRM por lote",10,200,50,step=10)
server=st.text_input("Servidor OSRM",value=OSRM)
balance=st.checkbox("Equilibrar equipes quando as distâncias OSRM forem competitivas",value=True)

clusters=preparar_grupos_50m(pending,50)
route_key=hashlib.sha256((signature+str(n_near)+server).encode()).hexdigest()
if st.session_state.get("nipv2_route_key")!=route_key:
    st.session_state["nipv2_candidates"]=candidatos_grupos(clusters,teams,n_near)
    st.session_state["nipv2_route_key"]=route_key
    st.session_state.pop("nipv2_zip",None)
candidates=st.session_state["nipv2_candidates"]
if not candidates.empty:
    remaining=candidates.index[candidates["STATUS_OSRM"].eq("PENDENTE")].tolist()
    completed=len(candidates)-len(remaining)
    st.progress(completed/max(1,len(candidates)),text=f"Consultas OSRM concluídas: {completed}/{len(candidates)}")
    st.caption(f"Restantes: {len(remaining)}. Consultas são realizadas somente ao clicar no botão, para evitar sobrecarregar o OSRM público.")
    if remaining and st.button("🛣️ Consultar próximo lote OSRM",type="primary"):
        start=time.monotonic()
        todo=remaining[:batch]
        progress=st.progress(0.0)
        for i,idx in enumerate(todo,1):
            r=candidates.loc[idx]
            answer=osrm_route(round(float(r["LAT_EQUIPE"]),6),round(float(r["LON_EQUIPE"]),6),
                              round(float(r["LAT_OBRA"]),6),round(float(r["LON_OBRA"]),6),server)
            candidates.at[idx,"STATUS_OSRM"]="CONFIRMADA" if answer.get("ok") else "FALHOU"
            if answer.get("ok"):
                candidates.at[idx,"KM_OSRM"]=answer["km"]
                candidates.at[idx,"MIN_OSRM"]=answer["min"]
            progress.progress(i/len(todo),text=f"OSRM {i}/{len(todo)} | {time.monotonic()-start:.1f}s")
        st.session_state["nipv2_candidates"]=candidates
        st.session_state.pop("nipv2_zip",None)
        st.rerun()
    with st.expander("Conferir consultas e distâncias"):
        st.dataframe(candidates,hide_index=True,use_container_width=True)
else:
    st.warning("Não há equipes elegíveis com coordenadas para os grupos encontrados.")

assigned=atribuir_por_osrm(clusters,candidates,teams,limit,balance,limite_km)
checks=validar_distribuicao_nova(pending,assigned,limit,limite_km)
if any(checks[k] for k in ["Duplicações indevidas","Notas ausentes","Equipes acima do limite","Equipes acima do limite KM","Grupos 50 m divididos"]):
    st.error("A validação encontrou divergências. Não exporte antes de corrigir.")
else:
    st.success("Contagem única, capacidade máxima e grupos de 50 m validados.")

st.subheader("2. Distribuição, equipes e notas sem atendimento")
a,b,c,d=st.columns(4)
for box,(label,key) in zip([a,b,c,d],[
    ("Notas pendentes","Notas originais"),("Alocadas com OSRM","Alocadas com OSRM"),
    ("Sem equipe confirmada","Sem alocação"),("Grupos de 50 m divididos","Grupos 50 m divididos")]):
    box.metric(label,checks[key])

load_panel=resumo_detalhado_equipes(assigned,teams)
if not load_panel.empty:
    visitas = assigned.loc[assigned["STATUS_DISTRIBUICAO"].eq("EQUIPE MAIS PROXIMA")].drop_duplicates(
        ["ID_EQUIPE", "GRUPO_50M"]
    )
    dist_por_equipe = visitas.groupby("ID_EQUIPE")["KM_IDA_VOLTA_GRUPO"].sum()
    load_panel["KM_ESTIMADOS_IDA_VOLTA"] = load_panel["ID_EQUIPE"].map(dist_por_equipe).fillna(0).round(2)
    load_panel["LIMITE_OBRAS"] = limit
    load_panel["LIMITE_KM"] = limite_km
    load_panel["OBRAS_RESTANTES"] = limit - load_panel["OBRAS"]
    load_panel["KM_RESTANTES_ESTIMADOS"] = (limite_km - load_panel["KM_ESTIMADOS_IDA_VOLTA"]).round(2)
st.dataframe(load_panel,hide_index=True,use_container_width=True)
if not load_panel.empty:
    st.bar_chart(load_panel.set_index("EQUIPE")["OBRAS"])

q=st.text_input("🔎 Consultar uma nota e a equipe mais próxima",placeholder="Digite o número da nota")
if q.strip():
    matched=assigned[assigned["NOTA"].astype(str).str.contains(re.escape(q.strip()),na=False)]
    if matched.empty:
        st.info("Nota não encontrada entre as pendentes. Verifique as excluídas na auditoria.")
    else:
        st.dataframe(matched[[x for x in ["NOTA","LISTA","MUNICIPIO","EQUIPE_MAIS_PROXIMA","BASE_ATRIBUIDA","DISTANCIA_EQUIPE_KM","RANKING_EQUIPES","STATUS_DISTRIBUICAO"] if x in matched]],hide_index=True,use_container_width=True)

without=assigned[assigned["STATUS_DISTRIBUICAO"]!="EQUIPE MAIS PROXIMA"].copy()
with st.expander(f"⚠️ Relatório detalhado de notas sem equipe ({len(without)})"):
    st.dataframe(without,hide_index=True,use_container_width=True)

st.subheader("3. Mapa interativo — cartões idênticos ao KML")
if st.checkbox("🗺️ Carregar mapa interativo",value=False):
    geo=assigned.dropna(subset=["LATITUDE","LONGITUDE"])
    if not geo.empty:
        mp=folium.Map(location=[float(geo["LATITUDE"].mean()),float(geo["LONGITUDE"].mean())],zoom_start=7,tiles="OpenStreetMap")
        folium.TileLayer("CartoDB positron",name="Mapa claro").add_to(mp)
        for cat in CATS:
            layer=folium.FeatureGroup(name=cat)
            cluster=MarkerCluster().add_to(layer)
            for _,r in geo[geo["LISTA"].eq(cat)].iterrows():
                folium.Marker([float(r["LATITUDE"]),float(r["LONGITUDE"])],
                              popup=folium.Popup(popup(r),max_width=355),
                              icon=folium.Icon(color=MAP_COLORS[cat])).add_to(cluster)
            layer.add_to(mp)
        team_layer=folium.FeatureGroup(name="Bases das equipes")
        for _,t in teams.iterrows():
            folium.Marker([float(t["LAT_EQUIPE"]),float(t["LON_EQUIPE"])],
                          tooltip=f"{t['EQUIPE']} — {t['TIPO_EQUIPE']}",
                          icon=folium.Icon(color="orange",icon="home")).add_to(team_layer)
        team_layer.add_to(mp)
        folium.LayerControl(collapsed=False).add_to(mp)
        st_folium(mp,height=610,use_container_width=True,key=f"nipv2_map_{versao}")
    else:
        st.warning("Nenhuma obra pendente tem coordenadas válidas.")

st.subheader("4. Exportações e auditoria")
params=pd.DataFrame([{"PARÂMETRO":k,"VALOR":str(v)} for k,v in {
    "Data/hora":datetime.now().isoformat(timespec="seconds"),"Arquivo saneamento":san_file.name,
    "Arquivo levantamento":lev_file.name,"Arquivo equipes":team_file.name,
    "Hash das três bases":signature,"Servidor OSRM":server,
    "Raio agrupamento metros":50,"Equipes candidatas":n_near,
    "Capacidade máxima de obras":limit,"Limite de km por equipe":limite_km,
    "Distancia extra habilitada":distancia_extra,
    "Calculo km": "Soma de 2 x OSRM base-grupo por visita/grupo de 50 m (nao rota sequencial)",
    "Equilibrar cargas":balance,
    "Método": "OSRM rodoviário confirmado; Haversine somente pré-filtro",
    "Garantia não cruzamento": "Não garantida; exige otimização de trajetos",
    **checks}.items()])
with st.expander("Parâmetros e auditoria de exclusões"):
    st.dataframe(params,hide_index=True,use_container_width=True)
    st.dataframe(audit,hide_index=True,use_container_width=True)

kml_general=kml_bytes(assigned.dropna(subset=["LATITUDE","LONGITUDE"]),name="NIP — Obras pendentes e equipes OSRM")
valid_export=not any(checks[k] for k in ["Duplicações indevidas","Notas ausentes","Equipes acima do limite","Equipes acima do limite KM","Grupos 50 m divididos"])

stamp=datetime.now().strftime("%d.%m.%Y")
preview=inventario_zips(assigned,teams,stamp)
with st.expander("📁 Prévia dos arquivos dos ZIPs"):
    st.dataframe(preview,hide_index=True,use_container_width=True)

zip_key=hashlib.sha256((assigned.to_json(orient="records",date_format="iso")+str(limit)+str(limite_km)+str(balance)).encode()).hexdigest()
if st.session_state.get("nipv2_zip_key")!=zip_key:
    st.session_state.pop("nipv2_zip",None)
    st.session_state["nipv2_zip_key"]=zip_key
if st.button("📦 Gerar ZIPs por equipe",disabled=not valid_export):
    bar=st.progress(0.0)
    timer=st.empty()
    start=time.monotonic()
    def progress_export(current,total,team):
        elapsed=time.monotonic()-start
        eta=elapsed/current*(total-current) if current else None
        bar.progress(current/max(1,total),text=f"Equipe {current}/{total}: {team}")
        timer.caption(f"Decorrido {elapsed:.1f}s | Restante estimado {eta:.1f}s" if eta is not None else f"Decorrido {elapsed:.1f}s")
    try:
        st.session_state["nipv2_zip"]=zip_distribuicao_osrm(assigned,teams,50,kml_general,progress_export)
        st.success(f"ZIPs concluídos em {time.monotonic()-start:.1f}s")
    except Exception as exc:
        st.error(f"Falha na geração: {exc}")

if st.session_state.get("nipv2_zip"):
    zex,zkm=st.session_state["nipv2_zip"]
    z1,z2=st.columns(2)
    z1.download_button("📥 ZIP Excel por equipe",zex,file_name=f"NIP_Equipes_Excel_{stamp}.zip",mime="application/zip")
    z2.download_button("📥 ZIP KML por equipe",zkm,file_name=f"NIP_Equipes_KML_{stamp}.zip",mime="application/zip")

x1,x2,x3=st.columns(3)
x1.download_button("🌍 KML geral",kml_general,file_name="NIP_Obras_Equipes_OSRM.kml",mime="application/vnd.google-earth.kml+xml")
x2.download_button("📊 Excel consolidado",excel_bytes({"PENDENTES":assigned,"EQUIPES":load_panel,"SEM EQUIPE":without,"OSRM":candidates,"AUDITORIA":audit,"PARAMETROS":params,"INCONSISTENCIAS":data["issues"]}),file_name="NIP_Analise_Completa.xlsx")
x3.download_button("📋 Excel FINL / CANC / Campo",excel_bytes({"FINL":finl,"CANC":canc,"CAMPO":field,"DUPLICADAS":duplicadas}),file_name="NIP_Exclusoes.xlsx")

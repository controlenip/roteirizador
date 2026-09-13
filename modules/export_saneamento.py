import io
import os
import re
import html
import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill, Font, Alignment

# ==========================================
# EXPORTADORES DO SANEAMENTO
# ==========================================

FORMULA_PREFIXES = ('=', '+', '-', '@')


def _safe_xml_id(value):
    text = re.sub(r'[^A-Za-z0-9_.-]+', '_', str(value or 'id')).strip('_.-')
    if not text:
        text = 'id'
    if not re.match(r'^[A-Za-z_]', text):
        text = f'id_{text}'
    return text


def _sanitize_excel_value(value):
    """Evita que texto vindo das bases seja interpretado como fórmula pelo Excel."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def sanitizar_dataframe_excel(df):
    out = df.copy()
    for col in out.columns:
        if str(out[col].dtype) == 'object':
            out[col] = out[col].map(_sanitize_excel_value)
    return out


def formatar_planilha_openpyxl(writer, sheet_name, max_linhas_largura=751):
    worksheet = writer.sheets[sheet_name]

    header_fill = PatternFill(start_color='002060', end_color='002060', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, name='Calibri')

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = 'A2'

    max_row = min(worksheet.max_row, int(max_linhas_largura))
    for col_cells in worksheet.iter_cols(min_row=1, max_row=max_row):
        max_length = 0
        column = col_cells[0].column_letter
        for cell in col_cells:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        worksheet.column_dimensions[column].width = min(max_length + 2, 60)


def gerar_excel_saneamento(df, colunas_originais=None):
    output = io.BytesIO()
    df_saida = df.loc[:, ~df.columns.duplicated()].copy()

    for col_name in ['PROTOCOLO', 'NOTA']:
        if col_name in df_saida.columns:
            df_saida = df_saida[~df_saida[col_name].isin(['RETORNO_BASE', 'PAUSA_ALMOCO'])]

    colunas_remover = [
        'ROTA_GEOMETRIA', '_HORA_INICIO_DT', '_HORA_FIM_DT', '_ORIGINAL_ROWS',
        '_ORIGEM_BASE', 'COR_ICONE', 'MUN_LIMPO', 'COORD_KEY', 'LAT_NUM', 'LON_NUM'
    ]
    df_saida = df_saida.drop(columns=[c for c in colunas_remover if c in df_saida.columns], errors='ignore')
    df_saida = sanitizar_dataframe_excel(df_saida)

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_saida.to_excel(writer, index=False, sheet_name='Saneamento Roteirizado')
        formatar_planilha_openpyxl(writer, 'Saneamento Roteirizado')

    return output.getvalue()


def gerar_excel_resumo_saneamento(df_resumo):
    output = io.BytesIO()
    df_resumo = df_resumo.loc[:, ~df_resumo.columns.duplicated()].copy()
    df_resumo = sanitizar_dataframe_excel(df_resumo)
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_resumo.to_excel(writer, index=False, sheet_name='Resumo Saneamento')
        formatar_planilha_openpyxl(writer, 'Resumo Saneamento')
    return output.getvalue()


def gerar_excel_generico(df, sheet_name='Dados'):
    output = io.BytesIO()
    df_saida = sanitizar_dataframe_excel(df.loc[:, ~df.columns.duplicated()].copy())
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if df_saida.empty:
            pd.DataFrame({'INFORMACAO': ['Nenhum registro']}).to_excel(writer, index=False, sheet_name=sheet_name[:31])
        else:
            df_saida.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        formatar_planilha_openpyxl(writer, sheet_name[:31])
    return output.getvalue()


def limpar_colunas_saneamento(df_alvo, cols_originais=None, colunas_selecionadas=None):
    """Monta o Excel operacional respeitando as colunas escolhidas na interface."""
    df_alvo = df_alvo.loc[:, ~df_alvo.columns.duplicated()].copy()

    if 'PROTOCOLO' in df_alvo.columns:
        if 'NOTA' in df_alvo.columns:
            df_alvo = df_alvo.drop(columns=['NOTA'])
        df_alvo = df_alvo.rename(columns={'PROTOCOLO': 'NOTA'})

    if 'BASE_ATRIBUIDA' in df_alvo.columns:
        if 'LEVANTADOR' in df_alvo.columns:
            df_alvo = df_alvo.drop(columns=['LEVANTADOR'])
        df_alvo = df_alvo.rename(columns={'BASE_ATRIBUIDA': 'LEVANTADOR'})

    obrigatorias = [
        'LEVANTADOR', 'NOME_DIA', 'DIA_MES', 'SEMANA', 'DIA', 'ORDEM',
        'HORA_INICIO', 'HORA_FIM', 'DISTANCIA_PONTO_ANTERIOR_KM',
        'DISTANCIA_ESTIMADA_KM', 'DISTANCIA_RODOVIARIA_KM', 'TEMPO_ROTA_MIN', 'TEMPO_ATENDIMENTO_MIN',
        'STATUS_ROTA', 'NOTA'
    ]

    if colunas_selecionadas:
        desejadas = []
        for c in colunas_selecionadas:
            cc = 'NOTA' if c == 'PROTOCOLO' else ('LEVANTADOR' if c == 'BASE_ATRIBUIDA' else c)
            if cc not in desejadas:
                desejadas.append(cc)
    else:
        desejadas = [
            'STATUS CLIENTE', 'NOME', 'TIPO DEMANDA', 'MUNICIPIO', 'ENDERECO',
            'BAIRRO', 'PONTO REFERENCIA', 'COMPLEMENTO', 'LATITUDE PROJETO',
            'LONGITUDE PROJETO', 'CLASSIFICACAO AREA', 'TEL FIXO', 'TEL MOVEL',
            'GRUPO TENSAO', 'ID', 'CONTA CONTRATO', 'EMPRESA', 'REGIONAL',
            'INSTALACAO', 'PRIORIDADE', 'SUPER_PONTO', 'ARQUIVO_ORIGEM',
            'QTD_OCORRENCIAS_NOTA', 'DUPLICADA_ENTRE_ARQUIVOS'
        ]

    final_cols = []
    for c in obrigatorias + desejadas:
        if c in df_alvo.columns and c not in final_cols:
            final_cols.append(c)
    return df_alvo[final_cols]


def gerar_kml_saneamento(df_kml, nome_arquivo, colunas_exibir, bases_ativas, tipo_periodo, funcao_formatadora):
    """KML com segmentos independentes; não cria ligações artificiais entre trechos."""
    df_kml = df_kml.copy()
    kml = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<kml xmlns="http://www.opengis.net/kml/2.2">', '<Document>',
        f'<name>{html.escape(str(nome_arquivo))}</name>'
    ]

    kml.append('<Style id="s_blue"><IconStyle><scale>1.2</scale><Icon><href>http://maps.google.com/mapfiles/kml/paddle/blu-blank.png</href></Icon></IconStyle></Style>')
    kml.append('<Style id="s_yellow"><IconStyle><scale>1.3</scale><Icon><href>http://maps.google.com/mapfiles/kml/paddle/ylw-blank.png</href></Icon></IconStyle></Style>')
    kml.append('<Style id="s_failed"><IconStyle><scale>1.2</scale><Icon><href>http://maps.google.com/mapfiles/kml/paddle/wht-blank.png</href></Icon></IconStyle></Style>')
    colors = ['ff0000ff', 'ff00ff00', 'ffff0000', 'ff00ffff', 'ffffff00', 'ffff00ff']

    for b_idx, b in enumerate(bases_ativas):
        if pd.isna(b) or b == 'NÃO ALOCADO':
            continue
        pasta = [f'<Folder><name>Equipe: {html.escape(str(b))}</name>']
        df_b = df_kml[df_kml['BASE_ATRIBUIDA'] == b]
        periodos = df_b['PERIODO'].dropna().unique() if 'PERIODO' in df_b.columns else [1]

        for p_idx, p in enumerate(periodos):
            df_p = df_b[df_b['PERIODO'] == p] if 'PERIODO' in df_b.columns else df_b
            c_line = colors[p_idx % len(colors)]
            style_id = _safe_xml_id(f's_line_{b_idx}_{p_idx}_{b}')
            kml.append(f'<Style id="{style_id}"><LineStyle><color>{c_line}</color><width>4</width></LineStyle></Style>')

            nome_pasta_per = f'{tipo_periodo} {p}' if tipo_periodo == 'Dia' else f'Semana {p}'
            pasta.append(f'<Folder><name>{html.escape(str(nome_pasta_per))}</name>')

            # MultiGeometry: cada deslocamento permanece um LineString independente.
            segmentos = []
            for _, r in df_p.iterrows():
                geom = r.get('ROTA_GEOMETRIA')
                if isinstance(geom, list) and len(geom) >= 2:
                    coords = []
                    for pt in geom:
                        if isinstance(pt, (list, tuple)) and len(pt) >= 2:
                            coords.append(f'{pt[0]},{pt[1]},0')
                    if len(coords) >= 2:
                        segmentos.append('<LineString><tessellate>1</tessellate><coordinates>\n' + '\n'.join(coords) + '\n</coordinates></LineString>')
            if segmentos:
                pasta.append(
                    f'<Placemark><name>Traçado da Rota</name><styleUrl>#{style_id}</styleUrl>'
                    '<MultiGeometry>' + ''.join(segmentos) + '</MultiGeometry></Placemark>'
                )

            for _, r in df_p.iterrows():
                if r.get('PROTOCOLO') in ['RETORNO_BASE', 'PAUSA_ALMOCO']:
                    continue
                lat, lon = r.get('LATITUDE'), r.get('LONGITUDE')
                if pd.isna(lat) or pd.isna(lon):
                    continue

                is_sp = isinstance(r.get('_ORIGINAL_ROWS'), list) and len(r['_ORIGINAL_ROWS']) > 1
                qty = len(r.get('_ORIGINAL_ROWS', [])) if is_sp else 1
                status_rota = str(r.get('STATUS_ROTA', ''))
                cor = 's_failed' if status_rota == 'SEM_ROTA_OSRM' else ('s_yellow' if is_sp else 's_blue')
                nome_ponto = f"[{r.get('ORDEM', '')}] 🏢 SUPER PONTO ({qty} un.)" if is_sp else f"[{r.get('ORDEM', '')}] Doc: {r.get('PROTOCOLO', '')}"
                bg_color = '#FFD700' if is_sp else ('#777777' if status_rota == 'SEM_ROTA_OSRM' else '#0D256C')
                text_color = '#000000' if is_sp else '#ffffff'

                desc = f'''<![CDATA[
                <div style="font-family:sans-serif; width:300px; border-radius:8px; overflow:hidden; box-shadow:0 2px 5px rgba(0,0,0,0.15);">
                    <div style="background:{bg_color}; color:{text_color}; padding:8px 10px; font-size:13px; font-weight:bold;">{html.escape(nome_ponto)}</div>
                    <div style="padding:10px; background:#fafafa; font-size:12px;">
                        <table style="width:100%; border-collapse:collapse;">
                '''

                extras = ['STATUS_ROTA', 'DISTANCIA_ESTIMADA_KM', 'DISTANCIA_RODOVIARIA_KM', 'TEMPO_ROTA_MIN', 'TEMPO_ATENDIMENTO_MIN', 'HORA_INICIO', 'HORA_FIM']
                cols_popup = list(dict.fromkeys(list(colunas_exibir or []) + extras))
                for c in cols_popup:
                    if c.upper() in ['PROTOCOLO', 'NOTA']:
                        continue
                    if c not in r.index:
                        continue
                    if is_sp and c.upper() not in ['LATITUDE', 'LONGITUDE', 'MUNICIPIO', 'LOCALIDADE', 'ZONA', 'REGIONAL'] and c in (colunas_exibir or []):
                        vals = [orig.get(c, '') for orig in r.get('_ORIGINAL_ROWS', [])]
                        val_html = "<div style='max-height:80px; overflow-y:auto; border:1px solid #ccc; padding:4px; background:#fff; border-radius:4px;'><ul style='margin:0; padding-left:0; list-style-type:none; font-size:11px; color:#333;'>" + ''.join([f"<li style='margin-bottom:2px;'><b>[{i+1}]</b> {funcao_formatadora(c, v)}</li>" for i, v in enumerate(vals)]) + '</ul></div>'
                    else:
                        val_html = funcao_formatadora(c, r.get(c, ''))
                    desc += f"<tr><td style='padding:3px 6px; font-weight:bold; color:#555; vertical-align:top; width:44%;'>{html.escape(str(c))}:</td><td style='padding:3px 6px; color:#333;'>{val_html}</td></tr>"

                link = f'https://www.google.com/maps?q={float(lat):.8f},{float(lon):.8f}'
                desc += f"</table><div style='margin-top:7px;'><a href='{html.escape(link, quote=True)}' target='_blank'>📍 Abrir no Google Maps</a></div></div></div>]]>"
                pasta.append(
                    f'<Placemark><name>{html.escape(nome_ponto)}</name><styleUrl>#{cor}</styleUrl>'
                    f'<description>{desc}</description><Point><coordinates>{lon},{lat},0</coordinates></Point></Placemark>'
                )

            pasta.append('</Folder>')
        pasta.append('</Folder>')
        kml.extend(pasta)

    kml.extend(['</Document>', '</kml>'])
    return '\n'.join(kml)


def gerar_gpx_simples(df_kml, nome_rota):
    """GPX com um trkseg por equipe+dia/período, incluindo o retorno à base."""
    df_kml = df_kml.copy()
    gpx = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<gpx version="1.1" creator="Roteirizador NIP" xmlns="http://www.topografix.com/GPX/1/1">',
        f'  <metadata><name>{html.escape(str(nome_rota))}</name></metadata>'
    ]

    for _, row in df_kml.iterrows():
        if row.get('PROTOCOLO') in ['RETORNO_BASE', 'PAUSA_ALMOCO']:
            continue
        lat, lon = row.get('LATITUDE'), row.get('LONGITUDE')
        if pd.notna(lat) and pd.notna(lon):
            gpx.append(f'  <wpt lat="{lat}" lon="{lon}"><name>{html.escape(str(row.get("PROTOCOLO", "Ponto")))}</name></wpt>')

    if 'ROTA_GEOMETRIA' in df_kml.columns and not df_kml.empty:
        group_cols = [c for c in ['BASE_ATRIBUIDA', 'DIA_MES', 'PERIODO'] if c in df_kml.columns]
        groups = df_kml.groupby(group_cols, sort=False, dropna=False) if group_cols else [(('Rota',), df_kml)]
        for chave, grp in groups:
            chave_t = chave if isinstance(chave, tuple) else (chave,)
            nome_seg = ' - '.join(str(x) for x in chave_t if pd.notna(x)) or 'Segmento'
            pontos = []
            for _, row in grp.iterrows():
                geom = row.get('ROTA_GEOMETRIA')
                if isinstance(geom, list):
                    for pt in geom:
                        if isinstance(pt, (list, tuple)) and len(pt) >= 2:
                            lon, lat = pt[0], pt[1]
                            pontos.append((lat, lon))
            if pontos:
                gpx.append(f'  <trk><name>{html.escape(str(nome_seg))}</name><trkseg>')
                for lat, lon in pontos:
                    gpx.append(f'      <trkpt lat="{lat}" lon="{lon}"></trkpt>')
                gpx.append('    </trkseg></trk>')

    gpx.append('</gpx>')
    return '\n'.join(gpx)


def injetar_logo():
    if os.path.exists('LOGO_NIP.png'):
        try:
            st.logo('LOGO_NIP.png', link='/')
        except Exception:
            st.logo('LOGO_NIP.png')


def identificar_icone_folium(row, colunas_disponiveis):
    if 'TIPO DEMANDA' in colunas_disponiveis:
        t = str(row.get('TIPO DEMANDA', '')).upper()
        if 'RELIGA' in t or 'LIGACAO' in t:
            return 'bolt'
    return 'map-marker'

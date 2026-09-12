import pandas as pd
import io
import html
import os
import streamlit as st
from openpyxl.styles import PatternFill, Font, Alignment


def _valor_alias(row, aliases, default='-'):
    for col in aliases:
        val = row.get(col)
        if pd.notna(val) and str(val).strip().lower() not in ['', 'nan', 'none']:
            return str(val).strip()
    return default


def gerar_excel_analise(dict_dfs):
    """Gera um arquivo Excel com múltiplas abas, preservando os dados e otimizando a formatação."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        escreveu_aba = False
        for sheet_name, df in dict_dfs.items():
            if df is None or df.empty:
                continue

            escreveu_aba = True
            df_saida = df.drop(
                columns=['_ORIGINAL_ROWS', 'LAT_NUM', 'LON_NUM', 'COR_MAPA', 'COR_NOME', 'CLUSTER_ID'],
                errors='ignore'
            ).copy()

            safe_sheet_name = str(sheet_name)[:31]
            df_saida.to_excel(writer, index=False, sheet_name=safe_sheet_name)

            worksheet = writer.sheets[safe_sheet_name]
            header_fill = PatternFill(start_color='002060', end_color='002060', fill_type='solid')
            header_font = Font(color='FFFFFF', bold=True, name='Calibri')

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.freeze_panes = 'A2'

            # Em bases grandes, medir todas as células para largura pode ser caro.
            # Amostra até 750 linhas sem alterar qualquer conteúdo exportado.
            max_linhas_largura = min(worksheet.max_row, 751)
            for col_cells in worksheet.iter_cols(min_row=1, max_row=max_linhas_largura):
                max_length = 0
                column = col_cells[0].column_letter
                for cell in col_cells:
                    try:
                        if cell.value is not None:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                worksheet.column_dimensions[column].width = min(max_length + 2, 60)

        # Proteção para cenários onde todos os dataframes recebidos estejam vazios.
        if not escreveu_aba:
            pd.DataFrame({'INFORMACAO': ['Nenhum registro para exportar']}).to_excel(
                writer, index=False, sheet_name='Sem Dados'
            )

    return output.getvalue()


def gerar_kml_analise(df):
    # Trabalha sempre em cópia para não alterar o DataFrame usado pela tela/exportações.
    df = df.copy()

    kml = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<kml xmlns="http://www.opengis.net/kml/2.2">',
        '<Document>',
        '<name>Análise Cruzada NIP</name>'
    ]

    styles = {
        'red': 'http://maps.google.com/mapfiles/kml/paddle/red-blank.png',
        'orange': 'http://maps.google.com/mapfiles/kml/paddle/ylw-blank.png',
        'green': 'http://maps.google.com/mapfiles/kml/paddle/grn-blank.png',
        'purple': 'http://maps.google.com/mapfiles/kml/paddle/purple-blank.png',
        'blue': 'http://maps.google.com/mapfiles/kml/paddle/blu-blank.png'
    }

    for color, url in styles.items():
        kml.append(
            f'<Style id="style_{color}"><IconStyle><scale>1.1</scale><Icon><href>{url}</href></Icon></IconStyle>'
            '<LabelStyle><scale>0</scale></LabelStyle></Style>'
        )

    kml.append(
        '<Style id="style_black"><IconStyle><scale>1.1</scale><color>ff000000</color>'
        '<Icon><href>http://maps.google.com/mapfiles/kml/paddle/wht-blank.png</href></Icon></IconStyle>'
        '<LabelStyle><scale>0</scale></LabelStyle></Style>'
    )

    if 'CLUSTER_ID' not in df.columns:
        df['CLUSTER_ID'] = range(len(df))

    clusters = []
    for cid, grp in df.groupby('CLUSTER_ID'):
        lat = str(grp['LATITUDE'].iloc[0]).strip()
        lon = str(grp['LONGITUDE'].iloc[0]).strip()
        if not lat or not lon or lat.lower() == 'nan' or lon.lower() == 'nan':
            continue

        c_names = grp['COR_NOME'].astype(str).tolist()

        if any('Preto' in c for c in c_names) or any('Inválidas' in c for c in c_names):
            c_nome = next(c for c in c_names if 'Inválidas' in c or 'Preto' in c)
        elif any('Vermelho' in c for c in c_names) or any('Duplicadas' in c for c in c_names):
            c_nome = next(c for c in c_names if 'Duplicadas' in c or 'Vermelho' in c)
        elif len(grp) > 1:
            c_nome = '🟠 Notas Próximas'
        else:
            c_nome = c_names[0]

        if 'Inválidas' in c_nome or 'Preto' in c_nome:
            c_mapa = 'black'
        elif 'Duplicadas' in c_nome or 'Vermelho' in c_nome:
            c_mapa = 'red'
        elif 'Próximas' in c_nome or 'Laranja' in c_nome:
            c_mapa = 'orange'
        else:
            c_mapa = str(grp['COR_MAPA'].iloc[0])
            if c_mapa not in styles and c_mapa != 'black':
                c_mapa = 'blue'

        titulo_card = f"📍 Obras no Local ({len(grp)})"
        desc = f'''<![CDATA[
        <div style="font-family:sans-serif; width:300px; max-height:360px; overflow-y:auto; border-radius:8px; overflow-x:hidden; box-shadow:0 2px 5px rgba(0,0,0,0.15);">
            <div style="background:#0D256C; color:#ffffff; padding:8px 10px; font-size:13px; font-weight:bold; text-align:center; position:sticky; top:0;">{titulo_card}</div>
            <div style="padding:10px; background:#fafafa; font-size:12px;">
        '''

        nomes_notas = []
        for _, r in grp.iterrows():
            n = html.escape(str(r.get('NOTA', '')))
            nomes_notas.append(n)
            mun = html.escape(str(r.get('MUNICIPIO', '')))
            o = html.escape(str(r.get('ORIGEM_BASE', '')))
            s = html.escape(str(r.get('SITUACAO SAP', '')))
            s_sisco = html.escape(_valor_alias(r, ['STATUS SISCO', 'STATUS_SISCO']))
            s_list = html.escape(_valor_alias(r, ['STATUS LIST', 'STATUS_LIST']))
            col = html.escape(str(r.get('COLABORADORES MAIS PROXIMOS', '')))
            dup = html.escape(str(r.get('DUPLICADA', '')))

            aviso_gps = ''
            if dup == 'SIM' and len(grp) == 1:
                aviso_gps = "<br><span style='color:red; font-size:10px;'>⚠️ A cópia desta nota está em outro ponto geográfico.</span>"

            desc += f'''
            <table style="width:100%; border-collapse:collapse; margin-bottom:8px;">
                <tr><td style="padding:2px;"><b>Nota:</b></td><td style="padding:2px;">{n}</td></tr>
                <tr><td style="padding:2px;"><b>Município:</b></td><td style="padding:2px;">{mun}</td></tr>
                <tr><td style="padding:2px;"><b>Origem:</b></td><td style="padding:2px;">{o}</td></tr>
                <tr><td style="padding:2px;"><b>Status SAP:</b></td><td style="padding:2px;">{s}</td></tr>
                <tr><td style="padding:2px;"><b>SISCO / LIST:</b></td><td style="padding:2px;">{s_sisco} / {s_list}</td></tr>
                <tr><td style="padding:2px;"><b>Equipes Perto:</b></td><td style="padding:2px;">{col}</td></tr>
                <tr><td style="padding:2px;"><b>Duplicada:</b></td><td style="padding:2px;">{dup}{aviso_gps}</td></tr>
            </table>
            <hr style="margin:4px 0; border:0; border-top:1px solid #ddd;">
            '''

        desc += '</div></div>]]>'

        # Mantém a nota no nome quando houver uma única obra. Em clusters grandes,
        # usa um nome curto e deixa a relação completa dentro do card.
        if len(grp) == 1:
            nome_placemark = nomes_notas[0] if nomes_notas else 'Obra'
        else:
            nome_placemark = f"{len(grp)} obras no local"

        clusters.append({
            'COR_NOME': c_nome,
            'COR_MAPA': c_mapa,
            'LAT': lat,
            'LON': lon,
            'DESC': desc,
            'NOME': nome_placemark
        })

    df_clusters = pd.DataFrame(clusters)

    if not df_clusters.empty:
        for nome_grupo, df_grupo in df_clusters.groupby('COR_NOME', sort=True):
            kml.append(f'<Folder><name>{html.escape(str(nome_grupo))}</name>')
            for _, r in df_grupo.iterrows():
                kml.append(
                    f'<Placemark><name>{html.escape(str(r["NOME"]))}</name>'
                    f'<styleUrl>#style_{r["COR_MAPA"]}</styleUrl>'
                    f'<description>{r["DESC"]}</description>'
                    f'<Point><coordinates>{r["LON"]},{r["LAT"]},0</coordinates></Point></Placemark>'
                )
            kml.append('</Folder>')

    kml.append('</Document></kml>')
    return "\n".join(kml)


def injetar_logo():
    if os.path.exists("LOGO_NIP.png"):
        st.logo("LOGO_NIP.png", icon_image=None)
